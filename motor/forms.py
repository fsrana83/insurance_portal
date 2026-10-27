import csv
import io
import re
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation

from django import forms
from django.contrib.auth.models import User
from django.forms import formset_factory

import json

from core import flex
from core.models import Currency, FlexField
from core.widgets import DateInput, IconRadioSelect, MoneyInput, ProductCardSelect, RiderCheckboxes

from .models import (
    FUEL_TYPES, USAGES, VEHICLE_TYPES, Cover, Customer, Driver, Endorsement, PolicyRisk, Product, ProductCover, Vehicle, Voucher, chassis_validator,
)


class StyledMixin:
    """Adds the design-system CSS classes to every widget."""

    def _style(self):
        for f in self.fields.values():
            w = f.widget
            if isinstance(w, (IconRadioSelect, ProductCardSelect, RiderCheckboxes, forms.CheckboxInput)):
                continue
            cls = "select" if isinstance(w, forms.Select) else "textarea" if isinstance(w, forms.Textarea) else "input"
            w.attrs["class"] = (w.attrs.get("class", "") + " " + cls).strip()


class CustomerForm(StyledMixin, forms.ModelForm):
    class Meta:
        model = Customer
        fields = ["customer_type", "name", "civil_id", "date_of_birth", "cr_number", "cr_registration_date", "cr_expiry_date", "nationality", "phone", "email", "address", "government_owned"]
        widgets = {"date_of_birth": DateInput(), "cr_registration_date": DateInput(), "cr_expiry_date": DateInput()}

    def clean(self):
        d = super().clean()
        if d.get("customer_type") == Customer.CORPORATE:
            if not d.get("cr_number"):
                self.add_error("cr_number", "The commercial registration number is required for a fleet / corporate customer.")
            if not d.get("cr_registration_date"):
                self.add_error("cr_registration_date", "Enter the CR registration date.")
            if not d.get("cr_expiry_date"):
                self.add_error("cr_expiry_date", "Enter the CR expiry date.")
            elif d.get("cr_registration_date") and d["cr_expiry_date"] <= d["cr_registration_date"]:
                self.add_error("cr_expiry_date", "The CR expiry date must be after the registration date.")
        elif d.get("customer_type") == Customer.INDIVIDUAL:
            if not d.get("civil_id"):
                self.add_error("civil_id", "The ID card number is required for an individual.")
            if not d.get("date_of_birth"):
                self.add_error("date_of_birth", "Enter the date of birth.")
        from .customers import find_existing
        other = find_existing(d.get("civil_id") or "", d.get("cr_number") or "", exclude=self.instance)
        if other:
            self.add_error("civil_id" if d.get("civil_id") and other.civil_id == d.get("civil_id") else "cr_number",
                           f"This ID is already registered to {other.name}. A customer can have only one record.")
        return d

    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        flex.attach(self, FlexField.CUSTOMER, initial=self.instance.flex if self.instance.pk else None)
        self._style()

    def save(self, commit=True):
        obj = super().save(commit=False)
        obj.flex = flex.extract(self, FlexField.CUSTOMER)
        if commit:
            obj.save()
        return obj


class CurrencyChoice(forms.ModelChoiceField):
    def label_from_instance(self, obj):
        return f"{obj.code} — {obj.name}"


class QuoteForm(StyledMixin, forms.Form):
    # --- customer -----------------------------------------------------------
    customer = forms.ModelChoiceField(Customer.objects.all(), required=False, widget=forms.HiddenInput)
    new_customer_type = forms.ChoiceField(choices=Customer._meta.get_field("customer_type").choices, required=False, label="Type")
    new_name = forms.CharField(max_length=120, required=False, label="Full name / company")
    new_civil_id = forms.CharField(max_length=12, required=False, label="Civil ID / resident card")
    new_cr_number = forms.CharField(max_length=20, required=False, label="Commercial registration (CR)")
    new_dob = forms.DateField(required=False, widget=DateInput(), label="Date of birth")
    new_cr_registration_date = forms.DateField(required=False, widget=DateInput(), label="CR registration date")
    new_cr_expiry_date = forms.DateField(required=False, widget=DateInput(), label="CR expiry date")
    producer = forms.ModelChoiceField(User.objects.none(), required=False, label="Broker / agent")
    new_phone = forms.CharField(max_length=30, required=False, label="Mobile")
    new_email = forms.EmailField(required=False, label="Email")
    # --- policy ---------------------------------------------------------------
    product = forms.ModelChoiceField(Product.objects.none(), widget=ProductCardSelect, empty_label=None)
    currency = CurrencyChoice(Currency.objects.none(), empty_label=None, label="Issuance currency")
    start_date = forms.DateField(widget=DateInput(), label="Inception")
    end_date = forms.DateField(widget=DateInput(), label="Expiry")
    ncd_years = forms.IntegerField(min_value=0, max_value=10, initial=0, label="Claim-free years (NCD)")
    voucher_code = forms.CharField(max_length=24, required=False, label="Discount voucher",
                                   widget=forms.TextInput(attrs={"placeholder": "Enter code", "autocomplete": "off",
                                                                 "style": "text-transform:uppercase"}))
    discount_pct = forms.DecimalField(min_value=0, max_value=60, required=False, initial=0, decimal_places=2,
                                      label="Additional discount %")
    notes = forms.CharField(widget=forms.Textarea(attrs={"rows": 2}), required=False)

    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self.fields["product"].queryset = Product.objects.filter(active=True)
        self.fields["product"].widget.descriptions = {str(p.pk): p.description for p in Product.objects.filter(active=True)}
        self.fields["currency"].queryset = Currency.objects.filter(active=True)
        self.fields["currency"].widget.attrs["data-currency-source"] = "1"
        self.fields["new_customer_type"].initial = Customer.INDIVIDUAL
        self.fields["producer"].queryset = User.objects.filter(is_active=True, groups__name__in=["Agent", "Broker", "Production Manager", "Underwriter"]).distinct().order_by("username")
        self.fields["producer"].label_from_instance = lambda u: f"{u.get_full_name() or u.get_username()} ({u.profile.get_producer_type_display()})"
        flex.attach(self, FlexField.CUSTOMER, prefix="cflex_")
        flex.attach(self, FlexField.POLICY, initial=kw.get("initial", {}).get("flex") if isinstance(kw.get("initial"), dict) else None, prefix="pflex_")
        if not self.is_bound:
            self.fields["currency"].initial = Currency.objects.filter(is_base=True).first()
            self.fields["start_date"].initial = date.today()
            self.fields["end_date"].initial = date.today() + timedelta(days=364)
            p = Product.objects.filter(active=True).first()
            self.fields["product"].initial = p
        self._style()

    def clean(self):
        data = super().clean()
        s, e = data.get("start_date"), data.get("end_date")
        if s and e:
            if e <= s:
                self.add_error("end_date", "Expiry must be after inception.")
            elif (e - s).days > 366:
                self.add_error("end_date", "Motor policies are issued for a maximum of one year.")
            if s < date.today():
                self.add_error("start_date", "Back-dated cover is not allowed.")
        if not data.get("customer"):
            if not data.get("new_name"):
                self.add_error("new_name", "Select an existing customer or enter the new customer's name.")
            if not (data.get("new_civil_id") or data.get("new_cr_number")):
                self.add_error("new_civil_id", "Civil ID (individual) or CR number (corporate) is required.")
            if not data.get("new_phone"):
                self.add_error("new_phone", "Mobile number is required for SMS delivery.")
            from .customers import find_existing
            other = find_existing(data.get("new_civil_id") or "", data.get("new_cr_number") or "")
            if other:
                self.add_error("new_civil_id" if data.get("new_civil_id") else "new_cr_number",
                               f"{other.name} is already registered with this ID. Search for the customer above and select them instead of creating a duplicate.")
            ct = data.get("new_customer_type") or Customer.INDIVIDUAL
            if ct == Customer.INDIVIDUAL and not data.get("new_dob"):
                self.add_error("new_dob", "Date of birth is required (drives the young-driver excess).")
            if ct == Customer.INDIVIDUAL and not data.get("new_civil_id"):
                self.add_error("new_civil_id", "The ID card number is required for an individual.")
            if ct == Customer.CORPORATE:
                if not data.get("new_cr_number"):
                    self.add_error("new_cr_number", "The CR number is required for a fleet / corporate customer.")
                if not data.get("new_cr_registration_date"):
                    self.add_error("new_cr_registration_date", "Enter the CR registration date.")
                if not data.get("new_cr_expiry_date"):
                    self.add_error("new_cr_expiry_date", "Enter the CR expiry date.")
                elif data.get("new_cr_registration_date") and data["new_cr_expiry_date"] <= data["new_cr_registration_date"]:
                    self.add_error("new_cr_expiry_date", "The CR expiry date must be after the registration date.")
        data["discount_pct"] = data.get("discount_pct") or Decimal("0")
        code = (data.get("voucher_code") or "").strip().upper()
        data["voucher_code"] = code
        data["voucher"] = None
        if code and data.get("product"):
            from .services import validate_voucher
            v, err = validate_voucher(code, data["product"], data.get("customer"))
            if err:
                self.add_error("voucher_code", err)
            data["voucher"] = v
        return data

    def get_customer(self):
        c = self.cleaned_data
        if c.get("customer"):
            return c["customer"]
        return Customer.objects.create(
            customer_type=c.get("new_customer_type") or Customer.INDIVIDUAL, name=c["new_name"], civil_id=c.get("new_civil_id", ""),
            cr_number=c.get("new_cr_number", ""), date_of_birth=c.get("new_dob"), cr_registration_date=c.get("new_cr_registration_date"),
            cr_expiry_date=c.get("new_cr_expiry_date"), phone=c.get("new_phone", ""), email=c.get("new_email", ""),
            flex=flex.extract(self, FlexField.CUSTOMER),
        )


class RiskForm(StyledMixin, forms.Form):
    vehicle_id = forms.IntegerField(required=False, widget=forms.HiddenInput)
    make = forms.CharField(max_length=40, widget=forms.TextInput(attrs={"list": "make-list", "autocomplete": "off"}))
    model = forms.CharField(max_length=40, widget=forms.TextInput(attrs={"list": "model-list", "autocomplete": "off"}))
    year = forms.IntegerField(min_value=1980, max_value=date.today().year + 1)
    plate_code = forms.CharField(max_length=4, required=False, label="Plate letters")
    plate_number = forms.CharField(max_length=10, required=False, label="Plate number")
    chassis_no = forms.CharField(max_length=17, label="Chassis / VIN", validators=[chassis_validator])
    engine_no = forms.CharField(max_length=30, label="Engine no.")
    color = forms.CharField(max_length=30)
    vehicle_type = forms.ChoiceField(choices=VEHICLE_TYPES, widget=IconRadioSelect(kind="vehicle"), initial="SALOON", label="Vehicle type")
    usage = forms.ChoiceField(choices=USAGES, widget=IconRadioSelect(kind="usage"), initial="PRIVATE", label="Usage")
    fuel_type = forms.ChoiceField(choices=FUEL_TYPES, initial="PETROL", required=False, label="Fuel type")
    cc = forms.IntegerField(min_value=50, max_value=12000, initial=1600, label="Engine CC")
    seats = forms.IntegerField(min_value=1, max_value=80, initial=5, label="Seats")
    sum_insured = forms.DecimalField(min_value=0, required=False, decimal_places=3, widget=MoneyInput, label="Sum insured")
    riders = forms.ModelMultipleChoiceField(Cover.objects.none(), required=False, widget=RiderCheckboxes, label="Riders & extensions")
    si_note = forms.CharField(max_length=160, required=False, widget=forms.HiddenInput)
    drivers_json = forms.CharField(required=False, widget=forms.HiddenInput, label="Additional drivers")
    waived = forms.ModelMultipleChoiceField(Cover.objects.none(), required=False, label="Waive mandatory rider")
    waiver_reason = forms.CharField(max_length=200, required=False)

    def __init__(self, *a, allow_waiver=False, with_flex=True, **kw):
        super().__init__(*a, **kw)
        if with_flex:
            init = kw.get("initial") or {}
            flex.attach(self, FlexField.VEHICLE, initial=init.get("flex"), prefix="vflex_")
        riders = Cover.objects.filter(kind="RIDER")
        self.fields["riders"].queryset = riders
        meta = {}
        for c in riders:
            mand = ProductCover.objects.filter(cover=c, mandatory=True)
            meta[str(c.pk)] = {
                "mandatory_for": ",".join(str(pc.product_id) for pc in mand),
                "regulatory": any(pc.regulatory for pc in mand),
                "hint": c.description,
            }
        self.fields["riders"].widget.meta = meta
        waivable = Cover.objects.filter(id__in=ProductCover.objects.filter(waivable=True, regulatory=False).values("cover_id"))
        self.fields["waived"].queryset = waivable
        if not allow_waiver:
            del self.fields["waived"]
            del self.fields["waiver_reason"]
        self._style()

    def clean_chassis_no(self):
        return self.cleaned_data["chassis_no"].upper().replace(" ", "")

    def clean(self):
        d = super().clean()
        from .rules import usage_error
        err = usage_error(d.get("vehicle_type"), d.get("usage"))
        if err:
            self.add_error("usage", err)
        return d

    def clean_plate_code(self):
        return self.cleaned_data.get("plate_code", "").upper()

    def clean_fuel_type(self):
        return self.cleaned_data.get("fuel_type") or "PETROL"

    def clean_drivers_json(self):
        raw = (self.cleaned_data.get("drivers_json") or "").strip()
        if not raw:
            return []
        try:
            rows = json.loads(raw)
            assert isinstance(rows, list)
        except Exception:
            raise forms.ValidationError("The additional-driver list could not be read.")
        out = []
        for i, r in enumerate(rows, 1):
            f = DriverForm(r)
            if not f.is_valid():
                raise forms.ValidationError(f"Driver {i}: " + "; ".join(f"{k}: {', '.join(v)}" for k, v in f.errors.items()))
            out.append(f.cleaned_data)
        return out

    def clean(self):
        d = super().clean()
        if d.get("waived") and not d.get("waiver_reason"):
            self.add_error("waiver_reason", "A reason code is required to waive a mandatory rider.")
        ch = d.get("chassis_no")
        if ch and d.get("make") and d.get("model") and d.get("year"):
            # a chassis number identifies one vehicle: MAKE + MODEL + YEAR + CHASSIS is unique
            other = Vehicle.objects.filter(chassis_no=ch).first()
            if other and (d.get("vehicle_id") != other.pk):
                if Vehicle.make_key(d["make"], d["model"], d["year"], ch) == other.vehicle_key:
                    pass                      # same vehicle: it will be re-used
                else:
                    self.add_error("chassis_no", f"Chassis {ch} is already registered as {other.make} {other.model} {other.year}. "
                                                 "Make, model, year and chassis number must always match.")
            elif other and d.get("vehicle_id") == other.pk:
                pass
            new_key = Vehicle.make_key(d["make"], d["model"], d["year"], ch)
            if Vehicle.objects.filter(vehicle_key=new_key).exclude(chassis_no=ch).exists():  # pragma: no cover - defensive
                self.add_error("chassis_no", "This make / model / year / chassis combination already exists.")
        return d


class DriverForm(forms.Form):
    name = forms.CharField(max_length=100)
    licence_no = forms.CharField(max_length=30, label="Driving licence no.")
    licence_issue_date = forms.DateField(required=False)
    date_of_birth = forms.DateField()
    nationality = forms.CharField(max_length=40, required=False)
    relationship = forms.ChoiceField(choices=Driver.RELATIONS, required=False)

    def clean_date_of_birth(self):
        d = self.cleaned_data["date_of_birth"]
        if d > date.today().replace(year=date.today().year - 17):
            raise forms.ValidationError("A driver must be at least 17 years old.")
        return d


class BaseRiskFormSet(forms.BaseFormSet):
    product = None

    def clean(self):
        super().clean()
        seen = set()
        live = [f for f in self.forms if f.cleaned_data and not f.cleaned_data.get("DELETE")]
        if not live:
            raise forms.ValidationError("Add at least one vehicle.")
        for f in live:
            ch = f.cleaned_data.get("chassis_no")
            if ch in seen:
                f.add_error("chassis_no", "Duplicate chassis number in this fleet.")
            seen.add(ch)
            if self.product and self.product.includes_od and not f.cleaned_data.get("sum_insured"):
                f.add_error("sum_insured", "Sum insured is required for comprehensive cover.")


RiskFormSet = formset_factory(RiskForm, formset=BaseRiskFormSet, extra=0, min_num=1, validate_min=True, can_delete=True)


# ---------------------------------------------------------------------------
# Fleet upload
# ---------------------------------------------------------------------------
FLEET_COLUMNS = ["make", "model", "year", "plate_code", "plate_number", "chassis_no", "engine_no", "color",
                 "vehicle_type", "usage", "cc", "seats", "sum_insured"]
FLEET_OPTIONAL = ["fuel_type"]
_TYPE_LOOKUP = {**{k.lower(): k for k, _ in VEHICLE_TYPES}, **{l.lower(): k for k, l in VEHICLE_TYPES}}
_USAGE_LOOKUP = {**{k.lower(): k for k, _ in USAGES}, **{l.lower(): k for k, l in USAGES}}


class FleetUploadForm(StyledMixin, forms.Form):
    file = forms.FileField(label="Excel or CSV file")

    def clean_file(self):
        f = self.cleaned_data["file"]
        if not f.name.lower().endswith((".xlsx", ".csv")):
            raise forms.ValidationError("Upload a .xlsx or .csv file using the template.")
        if f.size > 5 * 1024 * 1024:
            raise forms.ValidationError("File is larger than 5 MB.")
        return f


def _read_rows(f):
    if f.name.lower().endswith(".csv"):
        text = f.read().decode("utf-8-sig")
        return list(csv.reader(io.StringIO(text)))
    from openpyxl import load_workbook

    wb = load_workbook(f, read_only=True, data_only=True)
    ws = wb.active
    return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]


def parse_fleet_file(f):
    """Parse a fleet file. Returns (valid_rows, errors) where errors is a list of (row_number, message)."""
    rows = _read_rows(f)
    if not rows:
        return [], [(1, "The file is empty.")]
    header = [h.strip().lower().replace(" ", "_") for h in rows[0]]
    missing = [c for c in FLEET_COLUMNS if c not in header]
    if missing:
        return [], [(1, "Missing columns: " + ", ".join(missing))]
    idx = {c: header.index(c) for c in FLEET_COLUMNS + [o for o in FLEET_OPTIONAL if o in header]}
    valid, errors, seen = [], [], set()
    for n, raw in enumerate(rows[1:], start=2):
        if not any((c or "").strip() for c in raw):
            continue
        rec = {c: (raw[i].strip() if i < len(raw) and raw[i] is not None else "") for c, i in idx.items()}
        rec["vehicle_type"] = _TYPE_LOOKUP.get(rec["vehicle_type"].lower(), rec["vehicle_type"])
        rec["usage"] = _USAGE_LOOKUP.get(rec["usage"].lower(), rec["usage"])
        if "fuel_type" in rec:
            rec["fuel_type"] = rec["fuel_type"].strip().upper().replace("PATROL", "PETROL")
        for numeric in ("year", "cc", "seats"):
            try:
                rec[numeric] = str(int(float(rec[numeric])))
            except ValueError:
                pass
        form = RiskForm(data=rec, prefix=None, with_flex=False)
        form.fields["sum_insured"].required = False
        if not form.is_valid():
            for field, errs in form.errors.items():
                errors.append((n, f"{field}: {'; '.join(errs)}"))
            continue
        ch = form.cleaned_data["chassis_no"]
        if ch in seen:
            errors.append((n, f"chassis_no: duplicate of an earlier row ({ch})"))
            continue
        seen.add(ch)
        valid.append({k: (str(v) if isinstance(v, Decimal) else v) for k, v in form.cleaned_data.items()
                      if k in FLEET_COLUMNS})
    return valid, errors


def fleet_template_csv():
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(FLEET_COLUMNS)
    w.writerow(["Toyota", "Land Cruiser", 2023, "AB", "1234", "JTMHV05J904123456", "1GR1234567", "White", "SUV", "Private", 4600, 7, 24500])
    w.writerow(["Isuzu", "NPR", 2022, "CD", "5678", "JAANPR75LN7100123", "4HK1100123", "White", "TRUCK", "Commercial", 5200, 3, 14000])
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Endorsement & decisions
# ---------------------------------------------------------------------------
class EndorsementForm(StyledMixin, forms.Form):
    endorsement_type = forms.ChoiceField(choices=Endorsement.TYPES, label="Endorsement")
    risk = forms.ModelChoiceField(PolicyRisk.objects.none(), required=False, label="Vehicle", empty_label=None)
    effective_date = forms.DateField(widget=DateInput(), initial=date.today, label="Effective date")
    new_end_date = forms.DateField(widget=DateInput(), required=False, label="New expiry date")
    new_engine_no = forms.CharField(max_length=30, required=False)
    new_color = forms.CharField(max_length=30, required=False)
    new_usage = forms.ChoiceField(choices=USAGES, required=False, widget=IconRadioSelect(kind="usage"), label="New usage")
    accessories_value = forms.DecimalField(min_value=0, required=False, decimal_places=3, widget=MoneyInput, label="Accessories value")
    accessories_note = forms.CharField(max_length=120, required=False, label="Accessories (e.g. sound system, bull bar)")
    new_owner = forms.ModelChoiceField(Customer.objects.all(), required=False, label="New owner")
    new_plate_code = forms.CharField(max_length=4, required=False, label="New plate letters")
    new_plate_number = forms.CharField(max_length=10, required=False, label="New plate number")
    driver_name = forms.CharField(max_length=100, required=False, label="Driver name")
    driver_licence = forms.CharField(max_length=30, required=False, label="Driving licence no.")
    driver_dob = forms.DateField(widget=DateInput(), required=False, label="Date of birth")
    driver_nationality = forms.CharField(max_length=40, required=False, label="Nationality")
    driver_relationship = forms.ChoiceField(choices=Driver.RELATIONS, required=False, label="Relationship")

    def __init__(self, policy, *a, **kw):
        super().__init__(*a, **kw)
        self.policy = policy
        self.fields["risk"].queryset = policy.risks.filter(active=True)
        self.fields["risk"].initial = policy.risks.filter(active=True).first()
        self._style()

    def clean(self):
        d = super().clean()
        t = d.get("endorsement_type")
        need_risk = {Endorsement.TRANSFER, Endorsement.ENGINE, Endorsement.COLOR, Endorsement.USAGE, Endorsement.ADDITIONS,
                     Endorsement.REGNO, Endorsement.ADDDRIVER}
        if t in need_risk and not d.get("risk"):
            self.add_error("risk", "Select the vehicle.")
        req = {Endorsement.EXTENSION: "new_end_date", Endorsement.ENGINE: "new_engine_no", Endorsement.COLOR: "new_color",
               Endorsement.USAGE: "new_usage", Endorsement.ADDITIONS: "accessories_value", Endorsement.TRANSFER: "new_owner",
               Endorsement.REGNO: "new_plate_number"}
        f = req.get(t)
        if f and not d.get(f):
            self.add_error(f, "This field is required for the selected endorsement.")
        if t == Endorsement.ADDDRIVER:
            for fld in ("driver_name", "driver_licence", "driver_dob"):
                if not d.get(fld):
                    self.add_error(fld, "Required to add a driver.")
        if t == Endorsement.USAGE and d.get("risk") and d.get("new_usage") == d["risk"].usage:
            self.add_error("new_usage", "Usage is unchanged.")
        from .rules import government_error, usage_error
        if t == Endorsement.USAGE and d.get("risk") and d.get("new_usage"):
            for err in (usage_error(d["risk"].vehicle_type, d["new_usage"]), government_error(d["new_usage"], self.policy.customer)):
                if err:
                    self.add_error("new_usage", err)
        if t == Endorsement.TRANSFER and d.get("new_owner") and d.get("risk"):
            err = government_error(d["risk"].usage, d["new_owner"])
            if err:
                self.add_error("new_owner", err)
        if t == Endorsement.TRANSFER and d.get("new_owner") and d["new_owner"] == self.policy.customer:
            self.add_error("new_owner", "The new owner is already the policy holder.")
        return d

    def service_kwargs(self):
        d = self.cleaned_data
        t = d["endorsement_type"]
        kw = {}
        if t == Endorsement.EXTENSION:
            kw["new_end_date"] = d["new_end_date"]
        elif t == Endorsement.ENGINE:
            kw["new_engine_no"] = d["new_engine_no"].upper()
        elif t == Endorsement.COLOR:
            kw["new_color"] = d["new_color"]
        elif t == Endorsement.USAGE:
            kw["new_usage"] = d["new_usage"]
        elif t == Endorsement.ADDITIONS:
            kw["accessories_value"] = d["accessories_value"]
            kw["note"] = d.get("accessories_note", "")
        elif t == Endorsement.TRANSFER:
            kw["new_customer"] = d["new_owner"]
        elif t == Endorsement.REGNO:
            kw["new_plate_code"] = (d.get("new_plate_code") or "").upper()
            kw["new_plate_number"] = d["new_plate_number"]
        elif t == Endorsement.ADDDRIVER:
            kw.update(driver_name=d["driver_name"], driver_licence=d["driver_licence"], driver_dob=d["driver_dob"],
                      driver_nationality=d.get("driver_nationality", ""), driver_relationship=d.get("driver_relationship") or "OTHER")
        return kw


class DecisionForm(StyledMixin, forms.Form):
    note = forms.CharField(max_length=300, required=False, widget=forms.TextInput(attrs={"placeholder": "Reason / note"}))

    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self._style()



# ---------------------------------------------------------------------------
# Vouchers, flex fields, company
# ---------------------------------------------------------------------------
class VoucherForm(StyledMixin, forms.ModelForm):
    count = forms.IntegerField(min_value=1, max_value=500, initial=1, label="How many codes", help_text="More than 1 generates unique codes from the prefix")
    prefix = forms.CharField(max_length=10, required=False, initial="PROMO", label="Code prefix")

    class Meta:
        model = Voucher
        fields = ["code", "description", "discount_type", "value", "currency", "valid_from", "valid_to", "max_uses",
                  "min_net_premium", "products", "customer", "active"]
        widgets = {"valid_from": DateInput(), "valid_to": DateInput(), "products": forms.CheckboxSelectMultiple()}

    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self.fields["code"].required = False
        self.fields["code"].help_text = "Leave blank to generate automatically"
        self.fields["products"].queryset = Product.objects.filter(active=True, is_tariff_regulated=False)
        self.fields["currency"].queryset = Currency.objects.filter(active=True)
        self.fields["customer"].queryset = Customer.objects.all()
        if self.instance.pk:
            del self.fields["count"], self.fields["prefix"]
        self._style()

    def clean(self):
        d = super().clean()
        if d.get("valid_from") and d.get("valid_to") and d["valid_to"] < d["valid_from"]:
            self.add_error("valid_to", "The end date is before the start date.")
        if d.get("discount_type") == Voucher.PERCENT and d.get("value") and not (0 < d["value"] <= 60):
            self.add_error("value", "A percentage voucher must be between 0 and 60.")
        if d.get("code") and not self.instance.pk and Voucher.objects.filter(code=d["code"].strip().upper()).exists():
            self.add_error("code", "This code already exists.")
        return d


class FlexFieldForm(StyledMixin, forms.ModelForm):
    class Meta:
        model = FlexField
        fields = ["entity", "label", "key", "field_type", "choices", "required", "help_text", "order", "active"]
        widgets = {"choices": forms.Textarea(attrs={"rows": 3})}

    def __init__(self, *a, **kw):
        super().__init__(*a, **kw)
        self.fields["key"].required = False
        self.fields["key"].help_text = "Leave blank to derive from the label"
        if self.instance.pk:
            self.fields["key"].disabled = True
            self.fields["entity"].disabled = True
        self._style()

    def clean(self):
        d = super().clean()
        from django.utils.text import slugify
        if not self.instance.pk:
            key = slugify(d.get("key") or d.get("label") or "").replace("-", "_")
            if not key:
                self.add_error("label", "Enter a label.")
            elif FlexField.objects.filter(entity=d.get("entity"), key=key).exists():
                self.add_error("key", "A flex field with this key already exists for that record type.")
            d["key"] = key
        if d.get("field_type") == FlexField.CHOICE and not (d.get("choices") or "").strip():
            self.add_error("choices", "List the drop-down options, one per line.")
        return d
