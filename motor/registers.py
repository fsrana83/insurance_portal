"""Claims registers and reinsurance-by-treaty reports.

Every figure comes from the IFRS 17 ledger (so the reports tie to the accounts) and is split by treaty with
``reinsurance.shares``. Amounts are in the base currency (OMR).

* Claim register          - claims registered in the period, position as at the end of the period
* Claims incurred         - incurred in the period = paid in the period + (OS at end - OS at start)
* Outstanding (OS) as at  - claims reserve as at a date
* Paid claims             - payments made in the period and reinsurance cash received
* Reinsurance by treaty   - premium ceded, claims incurred / paid / OS and recoveries per treaty
* Broker production       - production with commission earned, due, settled and balance
"""
import csv
from collections import OrderedDict
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.db.models import Min, Prefetch, Sum
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.contrib import messages

from accounting.models import JournalEntry, JournalLine
from accounting.services import _base
from claims.models import Claim, ClaimPayment
from core.models import user_role
from core.permissions import role_required, scope
from reinsurance import shares as ri
from reinsurance.models import Cession, Treaty

from . import commissions as comm
from .models import Commission, Policy

D = Decimal
ZERO = D("0")
HUNDRED = D(100)

CLAIM_ROLES = ("PRODUCTION_MANAGER", "UNDERWRITER", "FINANCE", "CLAIMS_MANAGER", "CLAIMS_HANDLER")
REPORTS = OrderedDict([
    ("claims-register", "Claim register"),
    ("claims-incurred", "Claims incurred"),
    ("claims-os", "Outstanding as at"),
    ("claims-paid", "Paid claims"),
    ("ri-treaty", "Reinsurance by treaty"),
    ("premium-treaty", "Premium by treaty"),
    ("treaty-recovery", "Treaty recovery journal"),
    ("injuries", "Injuries and deaths"),
    ("loss-ratio", "Loss ratio"),
    ("broker-commission", "Broker production & commission"),
    ("aging", "Receivables / payables ageing"),
])


# ------------------------------------------------------------------ helpers
def _date(v, default):
    try:
        return date.fromisoformat(v)
    except (TypeError, ValueError):
        return default


def _period(request):
    today = date.today()
    return _date(request.GET.get("from"), today.replace(day=1)), _date(request.GET.get("to"), today)


def gate(request, asat=False):
    """A report shows nothing until a period is chosen and Run is pressed. Returns (ran, error)."""
    g = request.GET
    if g.get("run") != "1":
        return False, ""
    try:
        if asat:
            date.fromisoformat(g.get("as_at", ""))
        else:
            a, b = date.fromisoformat(g.get("from", "")), date.fromisoformat(g.get("to", ""))
            if a > b:
                return False, "The start date is after the end date."
    except ValueError:
        return False, "Choose a valid " + ("date" if asat else "period (from and to)") + " first."
    return True, ""


def _claims(request):
    qs = Claim.objects.select_related("policy__customer", "policy__product", "policy__currency", "policy__producer", "risk__vehicle") \
        .prefetch_related(Prefetch("policy__cessions", queryset=Cession.objects.select_related("treaty", "fac_reinsurer")), "recoveries")
    return scope(request.user, qs, producer_field="policy__producer")


def _ledger(accounts, types=None, d_from=None, d_to=None):
    """{(claim_id, account): [debit, credit]} from the journal."""
    qs = JournalLine.objects.filter(entry__claim__isnull=False, account_id__in=accounts)
    if types:
        qs = qs.filter(entry__entry_type__in=types)
    if d_from:
        qs = qs.filter(entry__date__gte=d_from)
    if d_to:
        qs = qs.filter(entry__date__lte=d_to)
    out = {}
    for r in qs.values("entry__claim_id", "account_id").annotate(d=Sum("debit"), c=Sum("credit")):
        out[(r["entry__claim_id"], r["account_id"])] = [r["d"] or ZERO, r["c"] or ZERO]
    return out


def _dr(m, cid, acc):
    return m.get((cid, acc), [ZERO, ZERO])[0]


def _cr(m, cid, acc):
    return m.get((cid, acc), [ZERO, ZERO])[1]


def _registered(claim_ids=None):
    """Registration date of a claim = its first claim journal (the ledger carries the business date)."""
    q = JournalEntry.objects.filter(claim__isnull=False, entry_type="CLAIM").values("claim_id").annotate(d=Min("date"))
    return {r["claim_id"]: r["d"] for r in q}


class Table:
    def __init__(self, title, cols, rows, totals=(), note=""):
        Q = D("0.001")
        rows = [[v.quantize(Q) if isinstance(v, D) else v for v in r] for r in rows]
        self.title, self.cols, self.rows, self.note = title, cols, rows, note
        self.totals = None
        if rows and totals:
            self.totals = ["Total" if i == 0 else (sum((r[i] for r in rows), ZERO).quantize(Q) if i in totals else "") for i in range(len(cols))]

    @property
    def num_idx(self):
        return {i for i, (_, k) in enumerate(self.cols) if k in ("n", "i")}

    @property
    def int_idx(self):
        return {i for i, (_, k) in enumerate(self.cols) if k == "i"}


def _treaty_cols(per_row, keys_labels):
    """Order treaty keys by programme rank; FAC last."""
    rank = {t.reference: t.rank for t in Treaty.objects.all()}
    keys = sorted(keys_labels, key=lambda k: (k == "FAC", rank.get(k, 99), k))
    return keys


def _respond(request, name, tables, ctx):
    fmt = request.GET.get("format")
    if fmt == "csv":
        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = f'attachment; filename="{name}.csv"'
        w = csv.writer(resp)
        for t in tables:
            w.writerow([t.title])
            w.writerow([c for c, _ in t.cols])
            w.writerows([[str(v) if isinstance(v, date) else v for v in r] for r in t.rows])
            if t.totals:
                w.writerow(t.totals)
            w.writerow([])
        return resp
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        wb.remove(wb.active)
        for t in tables:
            import re as _re
            ws = wb.create_sheet(_re.sub(r"[\\/*?:\[\]]", "-", t.title)[:31 - len(str(len(wb.sheetnames)))] + (str(len(wb.sheetnames)) if any(t.title[:28] == n[:28] for n in wb.sheetnames) else ""))
            ws.append([c for c, _ in t.cols])
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F5C63")
            for r in t.rows:
                ws.append([float(v) if isinstance(v, D) else v for v in r])
            if t.totals:
                ws.append([float(v) if isinstance(v, D) else v for v in t.totals])
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(12, min(36, max(len(str(c.value or "")) for c in col[:60]) + 2))
            ws.freeze_panes = "A2"
        buf = BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
        return resp
    for t in tables:
        t.shown = t.rows[:1500]
    return render(request, "motor/register.html", {"tables": tables, "reports": REPORTS, **ctx})


def _shares(claim):
    return ri.claim_shares(claim)


# ------------------------------------------------------------------ reports
def _ri(claim, basis, xol_amount, loss_base, labels):
    """({key: amount}, total) of reinsurance on ``basis`` for a claim, registering treaty labels."""
    parts = ri.allocate(claim, basis, xol_amount, loss_base)
    labels.update({k: v[0] for k, v in parts.items()})
    d = {k: v[1] for k, v in parts.items()}
    return d, sum(d.values(), ZERO)


def _basis(m_claim, m_ri, cid):
    """(incurred per the claim journals, reinsurance asset booked) for a claim - both from the ledger."""
    inc = _dr(m_claim, cid, "5110") - _cr(m_claim, cid, "5110")
    return inc, _dr(m_ri, cid, "1230") - _cr(m_ri, cid, "1230")


def claims_register(request):
    s, e = _period(request)
    reg = _registered()
    claims = [c for c in _claims(request) if s <= reg.get(c.id, c.created_at.date()) <= e]
    m = _ledger(["2120"], d_to=e)
    cl = _ledger(["5110", "1230"], types=["CLAIM"], d_to=e)
    rec = _ledger(["5110"], types=["RECOVERY"], d_to=e)
    ri_cash = _ledger(["1230"], types=["RECOVERY"], d_to=e)
    paid = _ledger(["2120", "5110"], types=["CLAIM_PAY"], d_to=e)
    rows, per, labels = [], [], {}
    for c in claims:
        inc_c, ri_book = _basis(cl, cl, c.id)
        delay = _dr(paid, c.id, "5110")
        inc = inc_c + delay
        os_ = _cr(m, c.id, "2120") - _dr(m, c.id, "2120")
        paid_ = _dr(paid, c.id, "2120") + delay
        salvage = _cr(rec, c.id, "5110") - _dr(rec, c.id, "5110")
        tot = ri.total_share(_shares(c))
        parts, ri_inc = _ri(c, inc_c, ri_book - inc_c * tot, inc_c, labels)
        ri_recv = _cr(ri_cash, c.id, "1230")
        per.append(parts)
        rows.append([reg.get(c.id, c.created_at.date()), c.claim_no, c.date_of_loss, c.policy.policy_no, c.policy.customer.name, c.policy.product.name,
                     c.get_loss_type_display(), c.get_status_display(), inc, paid_, os_, salvage, ri_inc, ri_recv, ri_inc - ri_recv, inc - ri_inc - salvage])
    cols = [("Registered", "d"), ("Claim no.", "t"), ("Date of loss", "d"), ("Policy", "t"), ("Customer", "t"), ("Product", "t"), ("Loss type", "t"),
            ("Status", "t"), ("Incurred (OMR)", "n"), ("Paid (OMR)", "n"), ("Outstanding (OMR)", "n"), ("Salvage / subrogation (OMR)", "n"),
            ("RI recoverable (OMR)", "n"), ("RI received (OMR)", "n"), ("RI receivable (OMR)", "n"), ("Net retained (OMR)", "n")]
    keys = _treaty_cols(per, labels)
    for r, p in zip(rows, per):
        r.extend(p.get(k, ZERO) for k in keys)
    cols += [(f"RI recoverable: {labels[k]}", "n") for k in keys]
    rows.sort(key=lambda r: (r[0], r[1]))
    t = Table("Claim register", cols, rows, totals=set(range(8, len(cols))),
              note=f"Claims registered {s:%d %b %Y} to {e:%d %b %Y}; position as at {e:%d %b %Y}. Incurred is gross of reinsurance and net of the policy excess; "
                   "the treaty columns split the reinsurers' share (quota share, surplus and excess of loss) of incurred.")
    return [t], {"mode": "period", "date_from": s, "date_to": e}


def claims_incurred(request):
    s, e = _period(request)
    open_d = s - timedelta(days=1)
    reg = _registered()
    o = _ledger(["2120"], d_to=open_d)
    c_ = _ledger(["2120"], d_to=e)
    cl = _ledger(["5110", "1230"], types=["CLAIM"], d_from=s, d_to=e)
    cl_cum = _ledger(["5110"], types=["CLAIM"], d_to=e)
    paid = _ledger(["2120", "5110"], types=["CLAIM_PAY"], d_from=s, d_to=e)
    rec = _ledger(["5110"], types=["RECOVERY"], d_from=s, d_to=e)
    ri_cash = _ledger(["1230"], types=["RECOVERY"], d_from=s, d_to=e)
    rows, per, labels = [], [], {}
    for c in _claims(request):
        cid = c.id
        os_o = _cr(o, cid, "2120") - _dr(o, cid, "2120")
        os_c = _cr(c_, cid, "2120") - _dr(c_, cid, "2120")
        p_ = _dr(paid, cid, "2120") + _dr(paid, cid, "5110")
        inc = p_ + os_c - os_o
        inc_c, ri_book = _basis(cl, cl, cid)
        salv = _cr(rec, cid, "5110") - _dr(rec, cid, "5110")
        cash = _cr(ri_cash, cid, "1230")
        if not any((os_o, os_c, p_, inc, salv, cash, ri_book)):
            continue
        tot = ri.total_share(_shares(c))
        parts, ri_inc = _ri(c, inc_c, ri_book - inc_c * tot, _dr(cl_cum, cid, "5110") - _cr(cl_cum, cid, "5110"), labels)
        per.append(parts)
        rows.append([c.claim_no, reg.get(cid, c.created_at.date()), c.date_of_loss, c.policy.policy_no, c.policy.product.name, c.get_loss_type_display(),
                     c.get_status_display(), os_o, p_, os_c, inc, ri_inc, inc - ri_inc, salv, cash])
    cols = [("Claim no.", "t"), ("Registered", "d"), ("Date of loss", "d"), ("Policy", "t"), ("Product", "t"), ("Loss type", "t"), ("Status", "t"),
            ("OS at start (OMR)", "n"), ("Paid in period (OMR)", "n"), ("OS at end (OMR)", "n"), ("Incurred (OMR)", "n"),
            ("RI share of incurred (OMR)", "n"), ("Net incurred (OMR)", "n"), ("Salvage / subrogation received (OMR)", "n"), ("RI cash received (OMR)", "n")]
    keys = _treaty_cols(per, labels)
    for r, p in zip(rows, per):
        r.extend(p.get(k, ZERO) for k in keys)
    cols += [(f"RI incurred: {labels[k]}", "n") for k in keys]
    tables = [Table("Claims incurred", cols, rows, totals=set(range(7, len(cols))),
                    note=f"Incurred {s:%d %b %Y} to {e:%d %b %Y} = paid in the period + OS at {e:%d %b %Y} - OS at {open_d:%d %b %Y}. "
                         "Gross of reinsurance, net of policy excess.")]
    agg = OrderedDict()
    for r in rows:
        a = agg.setdefault(r[5], [0, ZERO, ZERO, ZERO, ZERO])
        a[0] += 1
        a[1] += r[8]
        a[2] += r[10]
        a[3] += r[11]
        a[4] += r[12]
    tables.append(Table("Incurred by loss type", [("Loss type", "t"), ("Claims", "i"), ("Paid (OMR)", "n"), ("Incurred (OMR)", "n"),
                                                  ("RI share (OMR)", "n"), ("Net (OMR)", "n")],
                        [[k, D(v[0]), v[1], v[2], v[3], v[4]] for k, v in agg.items()], totals={1, 2, 3, 4, 5}))
    return tables, {"mode": "period", "date_from": s, "date_to": e}


def claims_os(request):
    today = date.today()
    as_at = _date(request.GET.get("as_at"), today)
    m = _ledger(["2120"], d_to=as_at)
    cl = _ledger(["5110", "1230"], types=["CLAIM"], d_to=as_at)
    paid = _ledger(["2120", "5110"], types=["CLAIM_PAY"], d_to=as_at)
    reg = _registered()
    rows, per, labels = [], [], {}
    for c in _claims(request):
        os_ = _cr(m, c.id, "2120") - _dr(m, c.id, "2120")
        if os_ == 0:
            continue
        inc_c, ri_book = _basis(cl, cl, c.id)
        tot = ri.total_share(_shares(c))
        cum = (_dr(paid, c.id, "2120") + os_) or ZERO                      # paid + OS = cumulative incurred through the reserve
        f = os_ / cum if cum else ZERO
        parts, ri_os = _ri(c, os_, (ri_book - inc_c * tot) * f, inc_c, labels)
        r_date = reg.get(c.id, c.created_at.date())
        age = max((as_at - r_date).days, 0)
        bucket = "0-30" if age <= 30 else "31-90" if age <= 90 else "91-180" if age <= 180 else "181-365" if age <= 365 else "Over 1 year"
        per.append(parts)
        rows.append([c.claim_no, r_date, c.date_of_loss, c.policy.policy_no, c.policy.customer.name, c.policy.product.name, c.get_loss_type_display(),
                     c.get_status_display(), D(age), bucket, os_, ri_os, os_ - ri_os])
    cols = [("Claim no.", "t"), ("Registered", "d"), ("Date of loss", "d"), ("Policy", "t"), ("Customer", "t"), ("Product", "t"), ("Loss type", "t"),
            ("Status", "t"), ("Age (days)", "i"), ("Age band", "t"), ("OS (OMR)", "n"), ("RI share of OS (OMR)", "n"), ("Net OS (OMR)", "n")]
    keys = _treaty_cols(per, labels)
    for r, p in zip(rows, per):
        r.extend(p.get(k, ZERO) for k in keys)
    cols += [(f"RI OS: {labels[k]}", "n") for k in keys]
    rows.sort(key=lambda r: (r[1], r[0]))
    tables = [Table("Outstanding claims", cols, rows, totals=set(range(10, len(cols))),
                    note=f"Claims reserve as at {as_at:%d %b %Y}, from the ledger (liability for incurred claims, net of policy excess).")]
    bands = OrderedDict((b, [0, ZERO, ZERO]) for b in ["0-30", "31-90", "91-180", "181-365", "Over 1 year"])
    for r in rows:
        b = bands[r[9]]
        b[0] += 1
        b[1] += r[10]
        b[2] += r[11]
    tables.append(Table("Ageing of OS", [("Age band", "t"), ("Claims", "i"), ("OS (OMR)", "n"), ("RI share (OMR)", "n")],
                        [[k, D(v[0]), v[1], v[2]] for k, v in bands.items()], totals={1, 2, 3}))
    return tables, {"mode": "asat", "as_at": as_at}


def claims_paid(request):
    s, e = _period(request)
    claims = {c.id: c for c in _claims(request)}
    cl = _ledger(["5110", "1230"], types=["CLAIM"], d_to=e)
    c_ = _ledger(["2120"], d_to=e)
    paid_cum = _ledger(["2120"], types=["CLAIM_PAY"], d_to=e)
    pays = ClaimPayment.objects.filter(paid_on__gte=s, paid_on__lte=e, claim_id__in=list(claims)).select_related("claim").order_by("paid_on", "id")
    rows, per, labels = [], [], {}
    for p in pays:
        c = claims[p.claim_id]
        amt = _base(p.amount, c.policy)
        inc_c, ri_book = _basis(cl, cl, c.id)
        tot = ri.total_share(_shares(c))
        cum = _dr(paid_cum, c.id, "2120") + (_cr(c_, c.id, "2120") - _dr(c_, c.id, "2120"))
        f = amt / cum if cum else ZERO
        parts, ri_amt = _ri(c, amt, (ri_book - inc_c * tot) * f, inc_c, labels)
        per.append(parts)
        rows.append([p.paid_on, c.claim_no, c.policy.policy_no, c.policy.customer.name, c.policy.product.name, c.get_loss_type_display(),
                     p.get_stage_display(), p.payee, p.reference, c.policy.currency.code, p.amount, amt, ri_amt, amt - ri_amt])
    cols = [("Paid on", "d"), ("Claim no.", "t"), ("Policy", "t"), ("Customer", "t"), ("Product", "t"), ("Loss type", "t"), ("Stage", "t"),
            ("Payee", "t"), ("Reference", "t"), ("Currency", "t"), ("Amount", "n"), ("Paid (OMR)", "n"), ("RI share (OMR)", "n"), ("Net of RI (OMR)", "n")]
    keys = _treaty_cols(per, labels)
    for r, p_ in zip(rows, per):
        r.extend(p_.get(k, ZERO) for k in keys)
    cols += [(f"RI share: {labels[k]}", "n") for k in keys]
    idx = set(range(11, len(cols)))
    tables = [Table("Paid claims", cols, rows, totals=idx,
                    note=f"Payments made {s:%d %b %Y} to {e:%d %b %Y}. Reinsurers' share is what the treaties will be asked to reimburse.")]
    from claims.models import Recovery

    recs = Recovery.objects.filter(recovery_type="REINSURANCE", received_on__gte=s, received_on__lte=e, claim_id__in=list(claims)).order_by("received_on", "id")
    r2 = []
    for rc in recs:
        c = claims[rc.claim_id]
        inc_c, ri_book = _basis(cl, cl, c.id)
        tot = ri.total_share(_shares(c))
        parts, _t = _ri(c, inc_c, ri_book - inc_c * tot, inc_c, {})
        parts = {k: v for k, v in parts.items() if v}
        w = sum(parts.values(), ZERO)
        amt = _base(rc.amount, c.policy)
        names = {k: v[0] for k, v in ri.allocate(c, inc_c, ri_book - inc_c * tot, inc_c).items()}
        for k, v in parts.items():
            r2.append([rc.received_on, c.claim_no, c.policy.policy_no, names[k], amt * v / w if w else ZERO])
    tables.append(Table("Reinsurance recoveries received", [("Received on", "d"), ("Claim no.", "t"), ("Policy", "t"), ("Treaty", "t"), ("Received (OMR)", "n")],
                        r2, totals={4}, note="Cash received from reinsurers in the period, split by treaty in proportion to each treaty's share."))
    return tables, {"mode": "period", "date_from": s, "date_to": e}


def ri_treaty(request):
    s, e = _period(request)
    open_d = s - timedelta(days=1)
    from .reports_ext import _ceded_by_source

    booked = _ceded_by_source(s, e)
    pols = scope(request.user, Policy.objects.all()).prefetch_related(Prefetch("cessions", queryset=Cession.objects.select_related("treaty", "fac_reinsurer")))
    nums = {}
    for p in pols:
        for src in [p.policy_no] + [x.number for x in p.endorsements.all()]:
            if src in booked:
                nums[src] = p
    agg = OrderedDict()
    labels = {}

    def A(k):
        return agg.setdefault(k, {"prem": ZERO, "si": ZERO, "inc": ZERO, "paid": ZERO, "os": ZERO, "recv": ZERO})

    for src, p in nums.items():
        w = ri.policy_weights(p)
        for k, v in ri.split(booked[src], w).items():
            labels[k] = w[k][0]
            A(k)["prem"] += v
    for p in set(nums.values()):
        for c in p.cessions.all():
            k = ri.key_of(c)
            labels.setdefault(k, ri.label_of(c))
            A(k)["si"] += D(c.sum_insured_ceded)
    o = _ledger(["2120"], d_to=open_d)
    c_ = _ledger(["2120"], d_to=e)
    cl_w = _ledger(["5110", "1230"], types=["CLAIM"], d_from=s, d_to=e)
    cl_cum = _ledger(["5110"], types=["CLAIM"], d_to=e)
    paid = _ledger(["2120", "5110"], types=["CLAIM_PAY"], d_from=s, d_to=e)
    cash = _ledger(["1230"], types=["RECOVERY"], d_from=s, d_to=e)
    for c in _claims(request):
        cid = c.id
        os_c = _cr(c_, cid, "2120") - _dr(c_, cid, "2120")
        p_ = _dr(paid, cid, "2120") + _dr(paid, cid, "5110")
        inc_c, ri_book = _basis(cl_w, cl_w, cid)
        cs = _cr(cash, cid, "1230")
        if not any((os_c, p_, inc_c, ri_book, cs)):
            continue
        tot = ri.total_share(_shares(c))
        loss = _dr(cl_cum, cid, "5110") - _cr(cl_cum, cid, "5110")
        xol_w = ri_book - inc_c * tot
        inc_parts, _x = _ri(c, inc_c, xol_w, loss, labels)
        for k, v in inc_parts.items():
            A(k)["inc"] += v
        # paid / OS share: proportional treaties at their share, XoL pro-rata to the paid and OS fractions of the incurred movement
        for key, basis in (("paid", p_), ("os", os_c)):
            f = (basis / (p_ + os_c)) if (p_ + os_c) else ZERO
            parts, _y = _ri(c, basis, xol_w * f, loss, labels)
            for k, v in parts.items():
                A(k)[key] += v
        sh = {k: v for k, v in _ri(c, inc_c, xol_w, loss, labels)[0].items() if v}
        w = sum(sh.values(), ZERO)
        for k, v in sh.items():
            A(k)["recv"] += cs * v / w if w else ZERO
    rank = {t.reference: t.rank for t in Treaty.objects.all()}
    rows = []
    for k in sorted(agg, key=lambda k: (k == "FAC", rank.get(k, 99), k)):
        a = agg[k]
        rows.append([labels[k], a["si"], a["prem"], a["inc"], a["paid"], a["os"], a["recv"], a["paid"] - a["recv"]])
    cols = [("Treaty", "t"), ("Sum insured ceded, all policies (OMR)", "n"), ("Premium ceded in period (OMR)", "n"), ("Claims incurred share (OMR)", "n"),
            ("Paid share (OMR)", "n"), ("OS share at end (OMR)", "n"), ("Cash received (OMR)", "n"), ("Paid not yet recovered (OMR)", "n")]
    t = Table("Reinsurance by treaty", cols, rows, totals={2, 3, 4, 5, 6, 7},
              note=f"Period {s:%d %b %Y} to {e:%d %b %Y}. Proportional treaties follow the premium ceded per policy; excess-of-loss treaties show the premium "
                   "accrued (rate x net premium) and the recoveries posted with the claims. Sum insured is the current ceded exposure, not limited to the period.")
    return [t], {"mode": "period", "date_from": s, "date_to": e}


def premium_treaty(request):
    s, e = _period(request)
    from .reports_ext import _ceded_by_source

    booked = _ceded_by_source(s, e)
    pols = scope(request.user, Policy.objects.all()).select_related("product", "customer", "currency") \
        .prefetch_related(Prefetch("cessions", queryset=Cession.objects.select_related("treaty__reinsurer", "fac_reinsurer")), "endorsements")
    rows = []
    for p in pols:
        srcs = [(p.policy_no, p.issued_at.date(), "New business" if not p.proposal.renewal_of_id else "Renewal")] + \
               [(x.number, x.created_at.date(), f"Endorsement: {x.get_endorsement_type_display()}") for x in p.endorsements.all()]
        cess = sorted(p.cessions.all(), key=lambda c: (c.rank, c.id))
        w = ri.policy_weights(p)
        for src, d, typ in srcs:
            if src not in booked:
                continue
            sp = ri.split(booked[src], w)
            done = set()
            for c in cess:
                k = ri.key_of(c)
                if k in done:
                    continue
                done.add(k)
                grp = [x for x in cess if ri.key_of(x) == k]
                rein = (c.treaty.reinsurer.name if c.treaty_id else (c.fac_reinsurer.name if c.fac_reinsurer_id else "Unplaced"))
                ceded_prem = sum((D(x.premium_ceded) for x in grp), ZERO)
                pct = (sum((D(x.cession_pct) for x in grp), ZERO)) if len(grp) == 1 else (ceded_prem / D(p.net_premium * p.fx_rate) * HUNDRED if p.net_premium else ZERO)
                rows.append([d, src, typ, p.policy_no, p.customer.name, p.product.name, ri.label_of(c), c.get_treaty_type_display() if hasattr(c, "get_treaty_type_display") else c.treaty_type,
                             rein, pct, sum((D(x.sum_insured_ceded) for x in grp), ZERO) if src == p.policy_no else ZERO, sp.get(k, ZERO)])
    rows.sort(key=lambda r: (r[0], r[1], r[6]))
    cols = [("Date", "d"), ("Document", "t"), ("Type", "t"), ("Policy", "t"), ("Customer", "t"), ("Product", "t"), ("Treaty", "t"), ("Treaty type", "t"),
            ("Reinsurer", "t"), ("Cession %", "n"), ("Sum insured ceded (OMR)", "n"), ("Premium ceded (OMR)", "n")]
    return [Table("Premium by treaty", cols, rows, totals={10, 11},
                  note=f"Reinsurance distribution of every premium transaction {s:%d %b %Y} to {e:%d %b %Y}. Premium ceded is what the ledger booked "
                       "for the document, split across the policy's treaties.")], {"mode": "period", "date_from": s, "date_to": e}


def broker_commission(request):
    s, e = _period(request)
    producer = request.GET.get("producer") or None
    ptype = request.GET.get("ptype") or "BROKER"
    pols = scope(request.user, Policy.objects.filter(issued_at__date__gte=s, issued_at__date__lte=e)).select_related("producer", "product", "currency") \
        .prefetch_related("commissions", "payments")
    if producer:
        pols = pols.filter(producer_id=producer)
    if ptype != "ALL":
        pols = pols.filter(producer__profile__producer_type=ptype)
    agg = OrderedDict()
    lines = []
    for p in pols:
        ratio = comm.collected_ratio(p)
        net = p.net_premium * p.fx_rate
        total = p.total_payable * p.fx_rate
        got = total * ratio
        a = agg.setdefault(p.producer_id, {"name": p.producer.get_full_name() or p.producer.get_username(), "n": 0, "net": ZERO, "total": ZERO, "got": ZERO,
                                           "earned": ZERO, "due": ZERO, "paid": ZERO})
        a["n"] += 1
        a["net"] += net
        a["total"] += total
        a["got"] += got
        e_, d_, s_ = ZERO, ZERO, ZERO
        for c in p.commissions.all():
            e_ += D(c.amount)
            d_ += comm.eligible(c, ratio)
            s_ += D(c.settled_amount)
        a["earned"] += e_
        a["due"] += d_
        a["paid"] += s_
        lines.append([p.issued_at.date(), p.policy_no, a["name"], p.product.name, p.customer.name, net, got, ratio * HUNDRED, e_, d_, s_, d_ - s_])
    rows = []
    for a in agg.values():
        rows.append([a["name"], D(a["n"]), a["net"], a["total"], a["got"], (a["got"] / a["total"] * HUNDRED) if a["total"] else ZERO,
                     a["earned"], a["earned"] - a["due"], a["due"], a["paid"], a["due"] - a["paid"]])
    rows.sort(key=lambda r: -r[2])
    cols = [("Broker", "t"), ("Policies", "i"), ("Net premium (OMR)", "n"), ("Invoiced incl. VAT (OMR)", "n"), ("Premium collected (OMR)", "n"), ("Collected %", "n"),
            ("Commission earned (OMR)", "n"), ("Not yet due - premium uncollected (OMR)", "n"), ("Commission due (OMR)", "n"), ("Commission paid (OMR)", "n"),
            ("Balance payable (OMR)", "n")]
    tables = [Table("Broker production & commission", cols, rows, totals={1, 2, 3, 4, 6, 7, 8, 9, 10},
                    note=f"Business issued {s:%d %b %Y} to {e:%d %b %Y}. Earned = commission booked; Due = earned on premium collected so far (clawbacks deducted in full); "
                         "Balance = due - already paid.")]
    lines.sort(key=lambda r: (r[2], r[0], r[1]))
    tables.append(Table("Policy detail", [("Issued", "d"), ("Policy", "t"), ("Broker", "t"), ("Product", "t"), ("Customer", "t"), ("Net premium (OMR)", "n"),
                                          ("Collected (OMR)", "n"), ("Collected %", "n"), ("Earned (OMR)", "n"), ("Due (OMR)", "n"), ("Paid (OMR)", "n"),
                                          ("Balance (OMR)", "n")], lines, totals={5, 6, 8, 9, 10, 11}))
    return tables, {"mode": "period", "date_from": s, "date_to": e, "commission": True, "ptype": ptype,
                    "producers": User.objects.filter(policies__isnull=False).distinct(), "f": request.GET}


def loss_ratio(request):
    from . import lossratio as L

    s, e = _period(request)
    group = request.GET.get("group") if request.GET.get("group") in dict(L.GROUPS) else "insured"
    only = request.GET.get("only_claims") == "1"
    rows = L.compute(request.user, s, e, group)
    if only:
        rows = [r for r in rows if r["claims"] or r["incurred"]]
    rows.sort(key=lambda r: (r["loss_ratio"] is None, -(r["loss_ratio"] or 0), -r["incurred"]))
    lead = {"insured": [("Insured", "t"), ("ID / CR", "t")], "vehicle": [("Vehicle", "t"), ("Plate", "t"), ("Chassis", "t"), ("Current owner", "t")],
            "both": [("Insured", "t"), ("ID / CR", "t"), ("Vehicle", "t"), ("Plate", "t"), ("Chassis", "t")]}[group]
    cols = lead + [("Policies", "i"), ("Claims", "i"), ("Written premium (OMR)", "n"), ("Earned premium (OMR)", "n"), ("Paid (OMR)", "n"),
                   ("Outstanding (OMR)", "n"), ("Incurred (OMR)", "n"), ("Loss ratio %", "n")]
    out = []
    for r in rows:
        if group == "insured":
            head = [r["insured"], r["insured_id"]]
        elif group == "vehicle":
            head = [r["vehicle"], r["plate"], r["chassis"], r["owner"]]
        else:
            head = [r["insured"], r["insured_id"], r["vehicle"], r["plate"], r["chassis"]]
        out.append(head + [D(r["n_policies"]), D(r["claims"]), r["written"], r["earned"], r["paid"], r["os"], r["incurred"],
                           (r["loss_ratio"].quantize(D("0.1")) if r["loss_ratio"] is not None else "n/a")])
    t = Table(f"Loss ratio {dict(L.GROUPS)[group].lower()}", cols, out, totals=set(range(len(lead), len(cols) - 1)),
              note=f"{s:%d %b %Y} to {e:%d %b %Y}. Loss ratio = incurred claims / earned premium (both from the ledger). Worst ratios first; "
                   "'n/a' means no premium was earned in the period.")
    tot = L.totals(rows)
    if t.totals:
        t.totals[-1] = tot["loss_ratio"].quantize(D("0.1")) if tot["loss_ratio"] is not None else "n/a"
    return [t], {"mode": "period", "date_from": s, "date_to": e, "lr": True, "groups": L.GROUPS, "group": group, "only_claims": only}


def treaty_recovery(request):
    """Reinsurers' share of every claim payment, posted treaty by treaty when the payment was made, with what has been received."""
    from claims.models import TreatyRecovery
    s, e = _period(request)
    claims = {c.id: c for c in _claims(request)}
    qs = TreatyRecovery.objects.filter(posted_on__gte=s, posted_on__lte=e, claim_id__in=list(claims)).select_related("claim", "payment").order_by("treaty_ref", "posted_on", "id")
    rows = [[r.journal_no, r.posted_on, r.treaty_label or r.treaty_ref, r.claim.claim_no, r.claim.policy.policy_no, (r.payment.get_stage_display() if r.payment_id else "Adjustment"), r.amount, r.received, r.due]
            for r in qs]
    cols = [("Journal", "t"), ("Date", "d"), ("Treaty", "t"), ("Claim", "t"), ("Policy", "t"), ("Payment", "t"), ("Reinsurers' share (OMR)", "n"), ("Received (OMR)", "n"), ("Due (OMR)", "n")]
    t = Table("Treaty recovery journal", cols, rows, totals={6, 7, 8},
              note=f"{s:%d %b %Y} to {e:%d %b %Y}. Posted automatically when each claim payment is made: the reinsurers' share by treaty (excess of loss pro rata). "
                   "Cash received from reinsurers is applied to the oldest balances first.")
    return [t], {"mode": "period", "date_from": s, "date_to": e}


def injuries(request):
    """Bodily injuries and deaths on claims registered in the period."""
    from claims.models import Injury
    s, e = _period(request)
    claims = {c.id: c for c in _claims(request)}
    reg = _registered()
    qs = Injury.objects.filter(claim_id__in=list(claims)).select_related("claim", "hospital").order_by("claim__claim_no", "id")
    rows = []
    for i in qs:
        c = claims[i.claim_id]
        if not (s <= reg.get(c.id, c.created_at.date()) <= e):
            continue
        rows.append([c.claim_no, c.date_of_loss, c.policy.policy_no, i.person_name, i.get_role_display(), i.get_severity_display(), i.date_of_death or "",
                     i.hospital.name if i.hospital_id else "", _base(i.compensation_reserve, c.policy), _base(i.medical_cost, c.policy)])
    cols = [("Claim", "t"), ("Date of loss", "d"), ("Policy", "t"), ("Person", "t"), ("Role", "t"), ("Severity", "t"), ("Date of death", "d"), ("Hospital", "t"),
            ("Compensation (OMR)", "n"), ("Medical (OMR)", "n")]
    t = Table("Injuries and deaths", cols, rows, totals={8, 9}, note=f"Claims registered {s:%d %b %Y} to {e:%d %b %Y}.")
    return [t], {"mode": "period", "date_from": s, "date_to": e}


def aging(request):
    from . import aging as A
    as_at = _date(request.GET.get("as_at"), date.today())
    sections = A.build(as_at)
    nb = len(A.BANDS)
    sumrows, tables, side_tot = [], [], {"Receivable": [ZERO] * nb, "Payable": [ZERO] * nb}
    for name, (side, rows) in sections.items():
        tot = [sum((r[5 + i] for r in rows), ZERO) for i in range(nb)]
        sumrows.append([name, side, *tot, sum(tot, ZERO), D(len(rows))])
        for i in range(nb):
            side_tot[side][i] += tot[i]
    for side in ("Receivable", "Payable"):
        sumrows.append([f"Total {side.lower()}", side, *side_tot[side], sum(side_tot[side], ZERO), ZERO])
    net = [a - b for a, b in zip(side_tot["Receivable"], side_tot["Payable"])]
    sumrows.append(["Net position (receivable less payable)", "Net", *net, sum(net, ZERO), ZERO])
    cols = [("Section", "t"), ("Side", "t")] + [(b, "n") for b, _ in A.BANDS] + [("Total (OMR)", "n"), ("Items", "i")]
    tables.append(Table("Summary", cols, sumrows, note=f"Ageing as at {as_at:%d %b %Y}: days past the due date. Premium is due on issue (or per instalment on credit); claims and reinsurance terms are 30 days."))
    for name, (side, rows) in sections.items():
        rows = sorted(rows, key=lambda r: (-int(r[4]), r[0]))
        tables.append(Table(f"Detail - {name.split(' (')[0]}", A._cols("Reference"), rows, totals=set(range(5, 6 + nb))))
    return tables, {"mode": "asat", "as_at": as_at}


_FUNCS = {"aging": aging, "treaty-recovery": treaty_recovery, "injuries": injuries, "loss-ratio": loss_ratio, "claims-register": claims_register, "claims-incurred": claims_incurred, "claims-os": claims_os, "claims-paid": claims_paid,
          "ri-treaty": ri_treaty, "premium-treaty": premium_treaty, "broker-commission": broker_commission}


@role_required("AGENT", "BROKER", "CREDIT_CONTROL", *CLAIM_ROLES)
def register_view(request, slug):
    if slug not in _FUNCS:
        return redirect("register", slug="claims-register")
    role = user_role(request.user)
    if role in ("AGENT", "BROKER") and slug != "broker-commission":
        from django.core.exceptions import PermissionDenied

        raise PermissionDenied
    asat = slug in ("claims-os", "aging")
    ran, err = gate(request, asat)
    base = {"slug": slug, "title": REPORTS[slug], "can_settle": role in ("ADMIN", "FINANCE"), "mode": "asat" if asat else "period", "f": request.GET}
    if slug == "broker-commission":
        base.update({"commission": True, "ptype": request.GET.get("ptype") or "BROKER", "producers": User.objects.filter(policies__isnull=False).distinct()})
    if slug == "loss-ratio":
        from . import lossratio as L
        base.update({"lr": True, "groups": L.GROUPS, "group": request.GET.get("group") or "insured", "only_claims": request.GET.get("only_claims") == "1"})
    if not ran:
        base.update({"tables": [], "reports": REPORTS, "waiting": True, "error": err,
                     "date_from": _date(request.GET.get("from"), None), "date_to": _date(request.GET.get("to"), None), "as_at": _date(request.GET.get("as_at"), None)})
        return render(request, "motor/register.html", base)
    tables, ctx = _FUNCS[slug](request)
    ctx.update(base | ctx)
    return _respond(request, slug.replace("-", "_") + "_" + date.today().isoformat(), tables, ctx)


@role_required("FINANCE")
def settle_commission(request):
    """Pay the commission that is due to a producer (POST)."""
    if request.method != "POST":
        return redirect("register", slug="broker-commission")
    producer = User.objects.filter(pk=request.POST.get("producer")).first()
    if not producer:
        messages.error(request, "Choose a broker.")
        return redirect("register", slug="broker-commission")
    st = comm.settle(producer, request.user, reference=request.POST.get("reference", ""), note="Commission settlement")
    if st:
        messages.success(request, f"{st.number}: {st.amount} OMR commission paid to {producer.get_username()}.")
    else:
        messages.info(request, "Nothing is due to this broker - commission is payable only on collected premium.")
    return redirect("register", slug="broker-commission")
