"""Premium register: one line per premium transaction (new policy, renewal, endorsement, cancellation)."""
import csv
from datetime import date
from decimal import Decimal

from django.contrib.auth.models import User
from django.db.models import Prefetch
from django.http import HttpResponse
from django.shortcuts import render

from core.models import Currency
from core.permissions import role_required, scope

from .models import Endorsement, Policy, Product
from reinsurance import shares as ri
from reinsurance.models import Cession

D = Decimal
ZERO = D("0")
COLUMNS = ["Date", "Document", "Type", "Policy no.", "Customer", "Product", "Producer", "Currency", "Net premium", "Fees", "VAT",
           "Total", "Net premium (OMR)", "Total (OMR)", "Voucher discount", "Commission (OMR)", "Ceded premium (OMR)", "Paid", "Balance", "Retained premium (OMR)"]


def _dates(request):
    today = date.today()
    try:
        s = date.fromisoformat(request.GET.get("from", ""))
    except ValueError:
        s = today.replace(day=1)
    try:
        e = date.fromisoformat(request.GET.get("to", ""))
    except ValueError:
        e = today
    return s, e


def _ceded_by_source(s, e):
    """Premium ceded booked per document (policy / endorsement number) from the ledger, so the register ties to IFRS 17."""
    from django.db.models import Sum

    from accounting.models import JournalLine

    qs = JournalLine.objects.filter(account_id="2230", entry__entry_type__in=["ISSUE", "ENDORSE"], entry__date__gte=s, entry__date__lte=e) \
        .values("entry__source").annotate(c=Sum("credit"), d=Sum("debit"))
    return {r["entry__source"]: (r["c"] or ZERO) - (r["d"] or ZERO) for r in qs}


def register_rows(request, s, e, producer=None, product=None, currency=None):
    """Returns (columns, rows). Treaty columns (ceded premium per treaty) are added after the fixed columns."""
    base = Currency.base()
    booked = _ceded_by_source(s, e)
    pols = scope(request.user, Policy.objects.filter(issued_at__date__gte=s, issued_at__date__lte=e)).select_related(
        "customer", "product", "producer", "currency", "proposal").prefetch_related(
        "commissions", "payments", Prefetch("cessions", queryset=Cession.objects.select_related("treaty", "fac_reinsurer")))
    ends = Endorsement.objects.filter(policy__in=scope(request.user, Policy.objects.all()), created_at__date__gte=s, created_at__date__lte=e) \
        .exclude(premium_adjustment=0, fee=0).select_related("policy__customer", "policy__product", "policy__producer", "policy__currency") \
        .prefetch_related(Prefetch("policy__cessions", queryset=Cession.objects.select_related("treaty", "fac_reinsurer")))
    if producer:
        pols, ends = pols.filter(producer_id=producer), ends.filter(policy__producer_id=producer)
    if product:
        pols, ends = pols.filter(product_id=product), ends.filter(policy__product_id=product)
    if currency:
        pols, ends = pols.filter(currency_id=currency), ends.filter(policy__currency_id=currency)
    rows, splits, labels = [], [], {}
    for p in pols:
        # a policy's figures at issuance = current figures minus the effect of its endorsements
        en = list(p.endorsements.all())
        net = p.net_premium - sum((x.premium_adjustment for x in en), ZERO)
        fees = p.fees - sum((x.fee for x in en), ZERO)
        tax = p.tax - sum((x.tax for x in en), ZERO)
        total = p.total_payable - sum((x.total for x in en), ZERO)
        fx = p.fx_rate
        comm = sum((c.amount for c in p.commissions.all() if c.kind in ("BASE", "OVERRIDER")), ZERO)
        w = ri.policy_weights(p)
        ceded = booked.get(p.policy_no, sum((c.premium_ceded for c in p.cessions.all()), ZERO)) if not en else booked.get(p.policy_no, ZERO)
        splits.append(ri.split(ceded, w))
        labels.update({k: v[0] for k, v in w.items()})
        paid = sum((x.amount for x in p.payments.all() if x.status == "PAID"), ZERO)
        rows.append([p.issued_at.date(), p.policy_no, "Renewal" if p.proposal.renewal_of_id else "New business", p.policy_no, p.customer.name,
                     p.product.name, p.producer.get_username(), p.currency.code, net, fees, tax, total, base.quantize(net * fx),
                     base.quantize(total * fx), p.voucher_discount, comm, ceded, paid, p.total_payable - paid,
                     base.quantize(net * fx) - ceded])
    for x in ends:
        p = x.policy
        fx = p.fx_rate
        w = ri.policy_weights(p)
        ceded = booked.get(x.number, ZERO)
        splits.append(ri.split(ceded, w))
        labels.update({k: v[0] for k, v in w.items()})
        rows.append([x.created_at.date(), x.number, "Cancellation" if x.endorsement_type == "CANCELLATION" else f"Endorsement: {x.get_endorsement_type_display()}",
                     p.policy_no, p.customer.name, p.product.name, p.producer.get_username(), p.currency.code, x.premium_adjustment, x.fee,
                     x.tax, x.total, base.quantize(x.premium_adjustment * fx), base.quantize(x.total * fx), ZERO,
                     sum((c.amount for c in p.commissions.all() if c.kind == "CLAWBACK"), ZERO) if x.endorsement_type == "CANCELLATION" else ZERO,
                     ceded, ZERO, ZERO, base.quantize(x.premium_adjustment * fx) - ceded])
    # treaty columns, ordered by position in the reinsurance programme
    rank = {t.reference: t.rank for t in __import__("reinsurance.models", fromlist=["Treaty"]).Treaty.objects.all()}
    keys = sorted(labels, key=lambda k: (k == "FAC", rank.get(k, 99), k))
    for r, sp in zip(rows, splits):
        r.extend(sp.get(k, ZERO) for k in keys)
    order = sorted(range(len(rows)), key=lambda i: (rows[i][0], rows[i][1]))
    rows = [rows[i] for i in order]
    return COLUMNS + [f"Ceded: {labels[k]}" for k in keys], rows


@role_required("AGENT", "BROKER", "PRODUCTION_MANAGER", "UNDERWRITER", "FINANCE", "CLAIMS_MANAGER")
def premium_register(request):
    from .registers import gate

    g = request.GET
    ran, err = gate(request)
    if not ran:
        def _d(v):
            try:
                return date.fromisoformat(v or "")
            except ValueError:
                return None
        return render(request, "motor/premium_register.html", {
            "waiting": True, "error": err, "cols": [], "rows": [], "n": 0, "totals": {}, "date_from": _d(g.get("from")), "date_to": _d(g.get("to")),
            "producers": User.objects.filter(policies__isnull=False).distinct(), "products": Product.objects.all(),
            "currencies": Currency.objects.filter(active=True), "f": g})
    s, e = _dates(request)
    cols, rows = register_rows(request, s, e, g.get("producer") or None, g.get("product") or None, g.get("currency") or None)
    fmt = g.get("format")
    if fmt == "csv":
        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = f'attachment; filename="premium_register_{s}_{e}.csv"'
        w = csv.writer(resp)
        w.writerow(cols)
        w.writerows(rows)
        return resp
    if fmt == "xlsx":
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Premium register"
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F5C63")
        for r in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(34, max(len(str(c.value or "")) for c in col[:60]) + 2))
        ws.freeze_panes = "A2"
        buf = BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="premium_register_{s}_{e}.xlsx"'
        return resp
    tot = {k: sum((r[i] for r in rows), ZERO) for k, i in (("net", 12), ("total", 13), ("comm", 15), ("ceded", 16), ("voucher", 14))}
    return render(request, "motor/premium_register.html", {
        "cols": cols, "rows": rows[:2000], "n": len(rows), "truncated": len(rows) > 2000, "totals": tot, "date_from": s, "date_to": e,
        "producers": User.objects.filter(policies__isnull=False).distinct(), "products": Product.objects.all(),
        "currencies": Currency.objects.filter(active=True), "f": g,
    })
