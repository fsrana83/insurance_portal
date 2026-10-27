"""Exports of the management insights pack: CSV, XLSX, PDF, DOCX, PPTX. All render from ``mgmt_pack.build()`` data."""
import csv
import io
from decimal import Decimal

D = Decimal

FORMATS = {
    "pdf": ("application/pdf", "pdf"),
    "pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "csv": ("text/csv", "csv"),
}


def fmt_kpi(k):
    v = k["value"]
    if v is None:
        return "n/a"
    if k["unit"] == "%":
        return f"{v:.1f}%"
    if k.get("int"):
        return f"{v:,.0f}"
    return f"{v:,.0f} {k['unit']}".strip()


def fmt_change(k):
    if k.get("change") is not None:
        return f"{k['change']:+.1f}%"
    if k.get("delta") is not None:
        return f"{k['delta']:+.1f} pts"
    return ""


def _n(v, dp=0):
    return "" if v is None else f"{v:,.{dp}f}"


def _p(v):
    return "n/a" if v is None else f"{v:.1f}%"


def tables(pack):
    """[(title, cols, rows)] - plain strings, used by csv / xlsx / docx / pdf."""
    cur = pack["currency"]
    out = [("Key indicators", ["Indicator", "This period", "Previous period", "Change"],
            [[k["label"], fmt_kpi(k), ("" if k["prev"] is None else (f"{k['prev']:.1f}%" if k["unit"] == "%" else f"{k['prev']:,.0f}")), fmt_change(k)] for k in pack["kpis"]]),
           ("Monthly trend", ["Month", f"Written premium ({cur})", f"Earned premium ({cur})", f"Incurred claims ({cur})", "Loss ratio"],
            [[r["label"], _n(r["gwp"]), _n(r["earned"]), _n(r["claims"]), _p(r["lr"])] for r in pack["trend"]])]
    head = ["Name", f"Written ({cur})", f"Earned ({cur})", f"Claims ({cur})", f"Commission ({cur})", f"Ceded ({cur})", f"RI share of claims ({cur})", "Loss ratio"]
    for title, key in (("Performance by product", "by_product"), ("Performance by producer", "by_producer")):
        out.append((title, head, [[r["name"], _n(r["gwp"]), _n(r["earned"]), _n(r["claims"]), _n(r["comm"]), _n(r["ceded"]), _n(r["ri"]), _p(r["lr"])] for r in pack[key]]))
    out.append(("Claims by cause of loss", ["Cause", "Claims registered", f"Incurred ({cur})", f"RI share ({cur})", f"Average ({cur})"],
                [[r["name"], r["count"], _n(r["claims"]), _n(r["ri"]), _n(r["avg"])] for r in pack["by_loss"]]))
    ri = pack["reinsurance"]
    out.append(("Reinsurance by treaty", ri["cols"], [[("" if v is None else (f"{v:,.0f}" if isinstance(v, (int, float, Decimal)) else v)) for v in r] for r in ri["rows"]]))
    if pack["outliers"]:
        out.append(("Findings needing review", ["Severity", "Finding", "Subject", "Reason"], [[o["severity"], o["title"], o["subject"], o["reason"]] for o in pack["outliers"]]))
    out.append(("Insights", ["Level", "Insight"], [[i["level"], i["text"]] for i in pack["insights"]]))
    return out


def _stem(pack):
    return f"management-insights-{pack['period']['start']}-to-{pack['period']['end']}"


def _title(pack):
    return f"Management insights - {pack['period']['start']:%d %b %Y} to {pack['period']['end']:%d %b %Y}"


def to_csv(pack):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([pack["company"], _title(pack)])
    for title, cols, rows in tables(pack):
        w.writerow([])
        w.writerow([title])
        w.writerow(cols)
        w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(pack):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for title, cols, rows in tables(pack):
        ws = wb.create_sheet(title[:31])
        ws.append([_title(pack)])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
        ws.append(cols)
        for c in ws[3]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F5C63")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        for i, col in enumerate(ws.columns, 1):
            width = max((len(str(c.value)) for c in col[2:] if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 70 if i > 1 else 40)
        ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(pack):
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, Spacer

    from core import pdfkit as k
    title = _title(pack)
    buf, doc = k.new_doc(title)
    story = [Paragraph(k.esc(pack["company"]), k.SMALL), Paragraph(k.esc(title), k.H1),
             Paragraph(f"Compared with {pack['period']['prev_start']:%d %b %Y} to {pack['period']['prev_end']:%d %b %Y}. Source: IFRS 17 ledger. Generated {pack['generated']:%d %b %Y}.", k.SMALL),
             Spacer(1, 4 * mm), Paragraph("Key findings", k.H2)]
    colour = {"good": "#1c7c3e", "warn": "#b36b00", "bad": "#b3261e", "info": "#0f5c63"}
    for i in pack["insights"]:
        story.append(Paragraph(f'<font color="{colour.get(i["level"], "#333")}">&#9679;</font> {k.esc(i["text"])}', k.BODY))
        story.append(Spacer(1, 1.5 * mm))
    for ttl, cols, rows in tables(pack):
        if ttl == "Insights" or not rows:
            continue
        story.append(Paragraph(k.esc(ttl), k.H2))
        data = [cols] + [[("" if v is None else str(v)) for v in r] for r in rows]
        n = len(cols)
        if ttl == "Findings needing review":
            wid = [18 * mm, 45 * mm, 30 * mm, 87 * mm]
        else:
            first = 46 * mm if n > 4 else 62 * mm
            wid = [first] + [(178 * mm - first) / (n - 1)] * (n - 1)
        story.append(k.grid(data, wid, align_right_cols=tuple(range(1, n)) if ttl != "Findings needing review" else (), font=7.5, zebra=True))
    k.render(buf, doc, story)
    return buf.getvalue()


def _chart_png(pack):
    """Trend chart as PNG for Word; None when matplotlib is not installed."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return None
    t = pack["trend"]
    if not t:
        return None
    fig, ax = plt.subplots(figsize=(7.2, 2.8), dpi=150)
    x = range(len(t))
    ax.bar([i - 0.2 for i in x], [float(r["gwp"]) for r in t], 0.4, label="Written", color="#0f5c63")
    ax.bar([i + 0.2 for i in x], [float(r["claims"]) for r in t], 0.4, label="Claims", color="#e9a23b")
    ax.set_xticks(list(x))
    ax.set_xticklabels([r["label"].split()[0] for r in t], fontsize=8)
    ax.legend(frameon=False, fontsize=8)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="y", labelsize=8)
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf


def to_docx(pack):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.8)
    d.add_paragraph(pack["company"]).runs[0].font.size = Pt(9)
    d.add_heading(_title(pack), 0)
    d.add_paragraph(f"Compared with {pack['period']['prev_start']:%d %b %Y} to {pack['period']['prev_end']:%d %b %Y}. Source: IFRS 17 ledger.")
    d.add_heading("Key findings", 1)
    colour = {"good": RGBColor(0x1c, 0x7c, 0x3e), "warn": RGBColor(0xb3, 0x6b, 0), "bad": RGBColor(0xb3, 0x26, 0x1e), "info": RGBColor(0x0f, 0x5c, 0x63)}
    for i in pack["insights"]:
        p = d.add_paragraph(style="List Bullet")
        r = p.add_run(i["text"])
        r.font.color.rgb = colour.get(i["level"], RGBColor(0, 0, 0))
    img = _chart_png(pack)
    if img:
        d.add_heading("Premium and claims by month", 1)
        d.add_picture(img, width=Cm(17))
    for ttl, cols, rows in tables(pack):
        if ttl == "Insights" or not rows:
            continue
        d.add_heading(ttl, 1)
        t = d.add_table(rows=1, cols=len(cols))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(cols):
            t.rows[0].cells[i].text = str(c)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = "" if v is None else str(v)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_pptx(pack):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Emu, Inches, Pt
    TEAL, DARK, AMBER = RGBColor(0x0f, 0x5c, 0x63), RGBColor(0x0a, 0x3d, 0x43), RGBColor(0xe9, 0xa2, 0x3b)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title, sub=""):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12), Inches(0.8)).text_frame
        tb.text = title
        tb.paragraphs[0].font.size, tb.paragraphs[0].font.bold, tb.paragraphs[0].font.color.rgb = Pt(28), True, DARK
        if sub:
            p = tb.add_paragraph(); p.text = sub; p.font.size = Pt(12); p.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        return s

    def text_box(s, x, y, w, h):
        tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap = True
        return tf

    p = pack["period"]
    s = slide("Management insights", f"{pack['company']} | {p['start']:%d %b %Y} to {p['end']:%d %b %Y}")
    tf = text_box(s, 0.6, 2.4, 12, 2)
    tf.text = f"Compared with {p['prev_start']:%d %b %Y} to {p['prev_end']:%d %b %Y}"
    tf.paragraphs[0].font.size = Pt(16)
    # KPI slides, 8 tiles each
    kp = pack["kpis"]
    for page in range(0, len(kp), 8):
        s = slide("Key indicators" if page == 0 else "Key indicators (continued)")
        for i, k in enumerate(kp[page:page + 8]):
            x, y = 0.6 + (i % 4) * 3.05, 1.7 + (i // 4) * 2.6
            box = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(2.85), Inches(2.3))
            box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xf3, 0xef, 0xe6); box.line.color.rgb = RGBColor(0xd8, 0xd2, 0xc4)
            tf = box.text_frame; tf.word_wrap = True
            tf.text = k["label"]; tf.paragraphs[0].font.size = Pt(12); tf.paragraphs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            q = tf.add_paragraph(); q.text = fmt_kpi(k); q.font.size = Pt(22); q.font.bold = True; q.font.color.rgb = DARK
            c = fmt_change(k)
            if c:
                q = tf.add_paragraph(); q.text = c + " vs previous"; q.font.size = Pt(12); q.font.color.rgb = TEAL
    # findings
    s = slide("Key findings")
    tf = text_box(s, 0.7, 1.5, 12, 5.5)
    for n, i in enumerate(pack["insights"][:9]):
        q = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        q.text = "• " + i["text"]; q.font.size = Pt(15); q.space_after = Pt(8)
    # trend chart
    if pack["trend"]:
        s = slide("Premium and claims by month")
        cd = CategoryChartData()
        cd.categories = [r["label"] for r in pack["trend"]]
        cd.add_series("Written premium", [float(r["gwp"]) for r in pack["trend"]])
        cd.add_series("Earned premium", [float(r["earned"]) for r in pack["trend"]])
        cd.add_series("Incurred claims", [float(r["claims"]) for r in pack["trend"]])
        ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.6), cd).chart
        ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
        for ser, col in zip(ch.plots[0].series, (TEAL, RGBColor(0xb9, 0xcf, 0xd1), AMBER)):
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
    # product loss ratio chart
    prod = [r for r in pack["by_product"] if r["lr"] is not None]
    if prod:
        s = slide("Loss ratio by product")
        cd = CategoryChartData()
        cd.categories = [r["name"] for r in prod]
        cd.add_series("Loss ratio %", [round(float(r["lr"]), 1) for r in prod])
        ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.6), cd).chart
        ch.has_legend = False
        ch.plots[0].has_data_labels = True
        ch.plots[0].series[0].format.fill.solid(); ch.plots[0].series[0].format.fill.fore_color.rgb = TEAL
    # tables
    for ttl, cols, rows in tables(pack):
        if ttl in ("Insights", "Key indicators", "Monthly trend") or not rows:
            continue
        for part in range(0, len(rows), 10):
            s = slide(ttl if part == 0 else f"{ttl} (continued)")
            chunk = rows[part:part + 10]
            gt = s.shapes.add_table(len(chunk) + 1, len(cols), Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.4 * (len(chunk) + 1))).table
            for j, c in enumerate(cols):
                gt.cell(0, j).text = str(c)
            for i, r in enumerate(chunk, 1):
                for j, v in enumerate(r):
                    gt.cell(i, j).text = "" if v is None else str(v)
            for row in gt.rows:
                for c in row.cells:
                    for q in c.text_frame.paragraphs:
                        q.font.size = Pt(10 if len(cols) < 7 else 8)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


BUILDERS = {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf, "docx": to_docx, "pptx": to_pptx}


def export(pack, fmt):
    mime, ext = FORMATS[fmt]
    return BUILDERS[fmt](pack), mime, f"{_stem(pack)}.{ext}"
