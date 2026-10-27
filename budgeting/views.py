import csv
from datetime import date
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from core.permissions import role_required
from motor.models import Product

from . import actuals as A, charts, engine, outliers
from .models import MONTHS, METRIC_SHORT, METRICS, BudgetLine, BudgetVersion

VIEW = ("FINANCE", "PRODUCTION_MANAGER", "UNDERWRITER", "CLAIMS_MANAGER")
EDIT = ("FINANCE", "PRODUCTION_MANAGER")
ALL_METRICS = [m for m, _, _ in METRICS] + list(("LOSS", "EXPENSE", "COMBINED"))


def _int(v, default):
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


def _cutoff(request):
    """?through=YYYY-MM names the last closed month; default = the previous calendar month."""
    t = request.GET.get("through", "")
    try:
        y, m = t.split("-")
        return int(y), int(m)
    except ValueError:
        return A.last_closed()


def _dec(v):
    try:
        return Decimal(str(v).replace(",", "").strip() or "0")
    except InvalidOperation:
        return None


# ---------------------------------------------------------------------------------------------- versions
@role_required(*VIEW)
def budget_list(request):
    versions = BudgetVersion.objects.select_related("parent", "created_by")
    years = sorted({v.year for v in versions}, reverse=True)
    return render(request, "budgeting/budget_list.html", {"versions": versions, "years": years, "can_edit": _can_edit(request)})


def _can_edit(request):
    from core.models import user_role
    return user_role(request.user) in EDIT + ("ADMIN",)


@role_required(*EDIT)
def budget_new(request):
    today = date.today()
    if request.method == "POST":
        P = request.POST
        year, kind, source = _int(P.get("year"), today.year + 1), P.get("kind", "BUDGET"), P.get("source", "BLANK")
        pct = _dec(P.get("pct")) or Decimal(0)
        src = BudgetVersion.objects.filter(pk=P.get("source_version") or 0).first()
        if source == "COPY" and not src:
            messages.error(request, "Choose the version to copy.")
        else:
            name = P.get("name") or f"{'Budget' if kind == 'BUDGET' else 'Forecast'} {year}"
            v = engine.create_version(name, year, kind, source, request.user, int(pct) if pct == int(pct) else float(pct), src, P.get("note", ""), normalise=bool(P.get("normalise")))
            messages.success(request, f"{v.name} created as a draft. Review and adjust the figures, then approve it.")
            return redirect("budget_edit", pk=v.pk)
    return render(request, "budgeting/budget_new.html", {"versions": BudgetVersion.objects.all(), "year": today.year + 1 if today.month > 6 else today.year})


@role_required(*VIEW)
def budget_edit(request, pk):
    v = get_object_or_404(BudgetVersion, pk=pk)
    products = list(Product.objects.all())
    if request.method == "POST":
        if not _can_edit(request):
            messages.error(request, "You cannot change budgets.")
            return redirect("budget_edit", pk=pk)
        if not v.editable:
            messages.error(request, "Only draft versions can be edited. Re-open it first.")
            return redirect("budget_edit", pk=pk)
        bad, n = 0, 0
        existing = {(l.product_id, l.metric, l.month): l for l in v.lines.all()}
        for p in products:
            for met, _, _ in METRICS:
                for m in range(1, 13):
                    raw = request.POST.get(f"v_{p.pk}_{met}_{m}")
                    if raw is None:
                        continue
                    amt = _dec(raw)
                    if amt is None or amt < 0:
                        bad += 1
                        continue
                    line = existing.get((p.pk, met, m))
                    if line is None:
                        BudgetLine.objects.create(version=v, product=p, metric=met, month=m, amount=amt)
                        n += 1
                    elif line.amount != amt:
                        line.amount = amt
                        line.save(update_fields=["amount"])
                        n += 1
        v.note = request.POST.get("note", v.note)[:200]
        v.name = request.POST.get("name", v.name)[:80] or v.name
        v.save(update_fields=["note", "name"])
        if bad:
            messages.warning(request, f"{bad} value(s) were not numbers (or negative) and were left unchanged.")
        messages.success(request, f"Saved ({n} figures changed).")
        return redirect("budget_edit", pk=pk)
    cube = engine.budget_cube(v)
    blocks = []
    for p in products:
        rows = []
        for met, label, _ in METRICS:
            vals = [cube.get((p.pk, met, m), Decimal(0)) for m in range(1, 13)]
            rows.append({"metric": met, "label": label, "cells": [(m, vals[m - 1]) for m in range(1, 13)], "total": sum(vals, Decimal(0))})
        blocks.append({"product": p, "rows": rows})
    return render(request, "budgeting/budget_grid.html", {"v": v, "blocks": blocks, "months": MONTHS, "can_edit": _can_edit(request) and v.editable,
                                                          "can_manage": _can_edit(request), "children": v.children.all(), "next_year": v.year + 1,
                                                          "methods": engine.METHODS, "today": date.today()})


@role_required(*EDIT)
@require_POST
def budget_action(request, pk, action):
    from core.models import user_role
    v = get_object_or_404(BudgetVersion, pk=pk)
    P = request.POST
    admin_like = user_role(request.user) in ("ADMIN", "FINANCE")
    if action == "approve" and v.status == v.DRAFT and admin_like:
        v.status = v.APPROVED
        v.save(update_fields=["status"])
        messages.success(request, f"{v.name} approved. It is now the comparison baseline for {v.year}.")
    elif action == "lock" and v.status == v.APPROVED and admin_like:
        v.status = v.LOCKED
        v.save(update_fields=["status"])
        messages.success(request, "Locked: figures can no longer be changed.")
    elif action == "reopen" and v.status != v.DRAFT and user_role(request.user) == "ADMIN":
        v.status = v.DRAFT
        v.save(update_fields=["status"])
        messages.warning(request, "Re-opened as a draft.")
    elif action == "delete" and v.status == v.DRAFT:
        v.delete()
        messages.success(request, "Draft deleted.")
        return redirect("budget_list")
    elif action == "rollover":
        pct = _dec(P.get("pct")) or Decimal(0)
        nv = engine.rollover(v, _int(P.get("year"), v.year + 1), P.get("basis", "VERSION"), float(pct), request.user, normalise=bool(P.get("normalise")))
        messages.success(request, f"{nv.name} created as a draft from {v.name}.")
        return redirect("budget_edit", pk=nv.pk)
    elif action == "forecast":
        try:
            y, m = P.get("through", "").split("-")
            cutoff = (int(y), int(m))
        except ValueError:
            cutoff = A.last_closed()
        fv = engine.build_forecast(v, v.year, P.get("method", "BUDGET_ADJ"), request.user, cutoff)
        messages.success(request, f"{fv.name} saved: actuals to the cut-off, projection after it.")
        return redirect("budget_edit", pk=fv.pk)
    else:
        messages.error(request, "That action is not available for this version.")
    return redirect("budget_edit", pk=pk)


# ---------------------------------------------------------------------------------------------- analytics
def _xlsx(title, header, rows):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F5C63")
    for r in rows:
        ws.append([float(x) if isinstance(x, Decimal) else x for x in r])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col[:60]) + 2))
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return resp


@role_required(*VIEW)
def analytics(request):
    g = request.GET
    today = date.today()
    year = _int(g.get("year"), today.year)
    mode = g.get("mode") if g.get("mode") in ("month", "quarter", "year") else "month"
    metric = g.get("metric") if g.get("metric") in ALL_METRICS else "GWP"
    product = _int(g.get("product"), 0) or None
    method = g.get("method") if g.get("method") in dict(engine.METHODS) else "BUDGET_ADJ"
    version = BudgetVersion.objects.filter(pk=g.get("version") or 0).first() or engine.active_budget(year)
    forecast = BudgetVersion.objects.filter(pk=g.get("forecast") or 0, kind="FORECAST").first()
    cutoff = _cutoff(request)
    versions_year = BudgetVersion.objects.filter(year=year)
    ctx = {"year": year, "years": sorted({v.year for v in BudgetVersion.objects.all()} | {today.year}, reverse=True), "mode": mode, "metric": metric,
           "product": product, "products": Product.objects.all(), "version": version, "forecast": forecast, "method": method, "methods": engine.METHODS,
           "budgets": versions_year.filter(kind="BUDGET"), "forecasts": versions_year.filter(kind="FORECAST"), "cutoff": cutoff,
           "through": f"{cutoff[0]}-{cutoff[1]:02d}", "metrics": [(m, METRIC_SHORT[m]) for m in ALL_METRICS], "metric_label": METRIC_SHORT[metric]}
    if not version:
        return render(request, "budgeting/analytics.html", ctx)
    res = engine.analysis(version, year, mode, metric, product, cutoff, forecast, method)
    rows = res["rows"]
    if g.get("format") == "xlsx":
        resp = _xlsx("Budget vs actual", ["Period", "Budget", "Budget to date", "Actual", "Variance", "Variance %", "Favourable", "Prior year", "YoY %", "Full-period outlook"],
                     [[r["label"], r["budget"], r["budget_to_date"], r["actual"], r["variance"], r["var_pct"], {True: "Yes", False: "No"}.get(r["fav"], ""), r["prior"], r["yoy"], r["outlook"]] for r in rows])
        resp["Content-Disposition"] = f'attachment; filename="budget_vs_actual_{year}_{metric}_{mode}.xlsx"'
        return resp
    if g.get("format") == "csv":
        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = f'attachment; filename="budget_vs_actual_{year}_{metric}_{mode}.csv"'
        w = csv.writer(resp)
        w.writerow(["Period", "Budget", "Budget to date", "Actual", "Variance", "Variance %", "Prior year", "YoY %", "Outlook"])
        for r in rows:
            w.writerow([r["label"], r["budget"], r["budget_to_date"], r["actual"], r["variance"], r["var_pct"], r["prior"], r["yoy"], r["outlook"]])
        return resp
    labels = [r["label"] for r in rows]
    unit = "%" if metric in ("LOSS", "EXPENSE", "COMBINED") else ""
    ctx.update(rows=rows, kpis=engine.kpis(version, year, product, cutoff, forecast, method),
               chart=charts.bars(labels, [("Budget", "ch-bud", [r["budget"] for r in rows], "bar"), ("Actual", "ch-act", [r["actual"] for r in rows], "bar"),
                                          ("Outlook", "ch-fc", [r["outlook"] for r in rows], "line")], title=f"{METRIC_SHORT[metric]} budget vs actual {year}", unit=unit),
               total=_totals(rows, metric))
    return render(request, "budgeting/analytics.html", ctx)


def _totals(rows, metric):
    if metric in ("LOSS", "EXPENSE", "COMBINED"):
        return None
    tb = sum((r["budget"] or 0 for r in rows), Decimal(0))
    tbd = sum((r["budget_to_date"] or 0 for r in rows), Decimal(0))
    ta = sum((r["actual"] or 0 for r in rows if r["actual"] is not None), Decimal(0))
    to = sum((r["outlook"] or 0 for r in rows), Decimal(0))
    tp = sum((r["prior"] or 0 for r in rows), Decimal(0))
    var = ta - tbd
    return {"budget": tb, "budget_to_date": tbd, "actual": ta, "variance": var, "var_pct": var / tbd * 100 if tbd else None, "outlook": to, "prior": tp,
            "fav": engine.favourable(metric, var)}


@role_required(*VIEW)
def forecast_view(request):
    g = request.GET
    method = g.get("method") if g.get("method") in dict(engine.METHODS) else "BUDGET_ADJ"
    cutoff = _cutoff(request)
    horizon = min(max(_int(g.get("horizon"), 12), 3), 24)
    rows = engine.rolling(cutoff, method, horizon)
    labels = [l for l, _ in rows]
    met = g.get("metric") if g.get("metric") in [m for m, _, _ in METRICS] else "GWP"
    # trailing actuals for context
    act = A.flat(cutoff[0] - 1, cutoff[0])
    hist = []
    for i in range(11, -1, -1):
        idx = cutoff[0] * 12 + cutoff[1] - 1 - i
        y, m = idx // 12, idx % 12 + 1
        hist.append((f"{MONTHS[m - 1]} {y}", {k: sum((act.get((p, k, y, m), Decimal(0)) for p in engine.pids()), Decimal(0)) for k in engine.STORED}))
    chart = charts.bars([l for l, _ in hist] + labels, [("Actual", "ch-act", [r[met] for _, r in hist] + [None] * len(rows), "bar"),
                                                         ("Forecast", "ch-fc-bar", [None] * len(hist) + [r[met] for _, r in rows], "bar")],
                        title=f"Rolling forecast: {METRIC_SHORT[met]}")
    rows_t = [{"label": l, **r} for l, r in rows]
    tot = {k: sum((r[k] for r in rows_t), Decimal(0)) for k in engine.STORED}
    if g.get("format") == "xlsx":
        resp = _xlsx("Rolling forecast", ["Month"] + [METRIC_SHORT[k] for k in engine.STORED], [[r["label"]] + [r[k] for k in engine.STORED] for r in rows_t])
        resp["Content-Disposition"] = 'attachment; filename="rolling_forecast.xlsx"'
        return resp
    return render(request, "budgeting/forecast.html", {"rows": rows_t, "tot": tot, "chart": chart, "method": method, "methods": engine.METHODS, "metric": met,
                                                       "metrics": [(m, METRIC_SHORT[m]) for m, _, _ in METRICS], "cutoff": cutoff, "through": f"{cutoff[0]}-{cutoff[1]:02d}",
                                                       "horizon": horizon, "stored": engine.STORED, "labels": METRIC_SHORT,
                                                       "saved": BudgetVersion.objects.filter(kind="FORECAST")[:10]})


@role_required(*VIEW)
def outlier_view(request):
    g = request.GET
    today = date.today()
    year = _int(g.get("year"), today.year)
    cutoff = _cutoff(request)

    def fl(k, default):
        try:
            return float(g.get(k, default))
        except ValueError:
            return default

    prm = {"z": fl("z", 3.0), "iqr": fl("iqr", 1.5), "min_var": fl("min_var", 15.0), "early_days": int(fl("early_days", 30)), "late_days": int(fl("late_days", 30))}
    try:
        d_from = date.fromisoformat(g.get("from", ""))
    except ValueError:
        d_from = date(year - 1, 1, 1) if g.get("scope") != "year" else date(year, 1, 1)
    try:
        d_to = date.fromisoformat(g.get("to", ""))
    except ValueError:
        d_to = today
    version = BudgetVersion.objects.filter(pk=g.get("version") or 0).first() or engine.active_budget(year)
    found = outliers.detect(version, year, cutoff, d_from, d_to, prm)
    kind = g.get("kind", "")
    shown = [o for o in found if not kind or o["kind"] == kind]
    counts = {k: sum(1 for o in found if o["kind"] == k) for k, _ in outliers.KINDS}
    if g.get("format") == "xlsx":
        resp = _xlsx("Outliers", ["Type", "Severity", "Finding", "Subject", "Value", "Expected", "Score", "Why"],
                     [[dict(outliers.KINDS)[o["kind"]], o["severity"], o["title"], o["subject"], o["value"], o["expected"], o["score"], o["reason"]] for o in shown])
        resp["Content-Disposition"] = 'attachment; filename="outliers.xlsx"'
        return resp
    return render(request, "budgeting/outliers.html", {"found": shown, "counts": counts, "kinds": outliers.KINDS, "kind": kind, "prm": prm, "year": year, "version": version,
                                                       "d_from": d_from, "d_to": d_to, "total": len(found), "high": sum(1 for o in found if o["severity"] == "high"),
                                                       "through": f"{cutoff[0]}-{cutoff[1]:02d}", "budgets": BudgetVersion.objects.filter(kind="BUDGET")})
